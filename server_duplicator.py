#!/usr/bin/env python3
"""
server_duplicator.py - Duplikasi Checksheet Langsung dari Server FactoryHub

Alur:
1. Login ke FactoryHub & buka template checksheet master part asal (misal: 76750B040P).
2. Scrape data inspeksi (seluruh item, standar, metode, toleransi) & unduh seluruh file gambar referensi dari server.
3. Buat folder lokal di documents/belum/<target_part>/ berisi:
   - File Excel checksheet (.xlsx) yang di-generate dari data server.
   - Folder images/ berisi file gambar referensi yang diunduh.
   - File info.json terhubung ke katalog FactoryHub.
4. Buka halaman pembuatan checksheet baru di FactoryHub untuk part target (misal: 76750B040PQ),
   upload gambar hasil download, dan isi tabel inspeksi.
"""

import os
import sys
import re
import json
import asyncio
import urllib.parse
from datetime import datetime
from typing import Dict, Any, List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from playwright.async_api import async_playwright, Page, BrowserContext

from catalog_manager import load_catalog, clean_code
from extractor import clear_document_caches
from automator import (
    login_factoryhub,
    find_existing_template,
    fill_checksheet_form,
    INDEX_URL,
    CREATE_URL
)
import logger


async def scrape_template_from_server(page: Page, source_part: str) -> Dict[str, Any]:
    """
    Cari template checksheet di FactoryHub untuk source_part dan ekstrak seluruh datanya.
    """
    clean_src = clean_code(source_part)
    print(f"[*] Mencari template checksheet di FactoryHub untuk part '{source_part}'...")

    template_info = await find_existing_template(page, source_part)
    edit_url = None
    if template_info:
        edit_url = template_info.get("edit_url") or template_info.get("editUrl")

    if not edit_url:
        # Cek dari catalog.json jika find_existing_template gagal
        catalog = load_catalog()
        for p in catalog.get("parts", []):
            if (p.get("clean_number") or clean_code(p.get("part_number", ""))) == clean_src:
                if p.get("has_template") and p.get("template_info"):
                    ti = p["template_info"]
                    edit_url = ti.get("edit_url") or ti.get("editUrl")
                    if edit_url:
                        print(f"[+] Ditemukan URL template dari cache katalog: {edit_url}")
                        break

    if not edit_url:
        raise FileNotFoundError(
            f"Template checksheet untuk part '{source_part}' tidak ditemukan di FactoryHub. "
            f"Pastikan nomor part asal sudah memiliki template checksheet di FactoryHub."
        )

    print(f"[+] Membuka halaman template: {edit_url}")
    await page.goto(edit_url, wait_until="networkidle")

    # Scrape form metadata, items, and images
    data = await page.evaluate(r"""() => {
        const docInp = document.querySelector('input[name="doc_number"]');
        const descInp = document.querySelector('input[name="description"]');
        const tbody = document.getElementById("inspection-tbody");
        const trs = tbody ? Array.from(tbody.querySelectorAll("tr.inspection-row, tr")) : [];

        const items = trs.map((tr, index) => {
            const getVal = (sel) => {
                const el = tr.querySelector(sel);
                return el ? el.value.trim() : '';
            };
            const itemNo = getVal('input[name*="[item_no]"]') || (index + 1).toString();
            const itemDesc = getVal('input[name*="[inspection_item]"]');
            const std = getVal('input[name*="[standard]"]');
            const method = getVal('input[name*="[method]"], input[name*="[instrument_tools]"]') ||
                           (tr.cells && tr.cells[5] && tr.cells[5].querySelector('input') ? tr.cells[5].querySelector('input').value.trim() : '');
            const master = getVal('input[name*="[master_data]"]') ||
                           (tr.cells && tr.cells[6] && tr.cells[6].querySelector('input') ? tr.cells[6].querySelector('input').value.trim() : '');

            return {
                item_no: itemNo,
                inspection_item: itemDesc,
                standard: std,
                method: method,
                master_data: master
            };
        }).filter(it => it.inspection_item || it.standard);

        // Cari URL gambar referensi yang tersimpan di template
        const imgUrls = [];
        const seen = new Set();
        const candidateImgs = Array.from(document.querySelectorAll('img[src*="/storage/checksheets/masters/"], a[href*="/storage/checksheets/masters/"]'));
        candidateImgs.forEach(el => {
            const url = el.src || el.href;
            if (url && url.includes('/storage/checksheets/masters/') && !seen.has(url)) {
                seen.add(url);
                imgUrls.push(url);
            }
        });

        const catSel = document.querySelector('select[name="category"]');
        const catVal = catSel ? catSel.value : 'Accuracy';

        return {
            doc_number: docInp ? docInp.value.trim() : '',
            description: descInp ? descInp.value.trim() : '',
            category: catVal,
            items: items,
            image_urls: imgUrls
        };
    }""")

    data["source_part"] = source_part
    data["edit_url"] = edit_url
    print(f"[+] Berhasil scrape data server: Doc='{data['doc_number']}', Desc='{data['description']}', "
          f"{len(data['items'])} titik inspeksi, {len(data['image_urls'])} file gambar.")
    return data


async def download_server_images(
    context: BrowserContext,
    image_urls: List[str],
    output_dir: str
) -> List[str]:
    """
    Unduh seluruh file gambar referensi dari server FactoryHub ke folder output_dir.
    """
    os.makedirs(output_dir, exist_ok=True)
    downloaded_paths = []

    print(f"[*] Mengunduh {len(image_urls)} file gambar referensi dari server FactoryHub...")
    for idx, url in enumerate(image_urls, 1):
        ext = ".png"
        if ".jpg" in url.lower() or ".jpeg" in url.lower():
            ext = ".jpg"
        elif ".webp" in url.lower():
            ext = ".webp"

        filename = f"sketch_{idx:02d}{ext}"
        dest_path = os.path.join(output_dir, filename)

        try:
            resp = await context.request.get(url, ignore_https_errors=True)
            if resp.status == 200:
                body = await resp.body()
                with open(dest_path, "wb") as f:
                    f.write(body)
                downloaded_paths.append(os.path.abspath(dest_path))
                print(f"    ✓ Diunduh ({idx}/{len(image_urls)}): {filename} ({len(body):,} bytes)")
            else:
                print(f"    [!] Gagal mengunduh gambar {url} (HTTP {resp.status})")
        except Exception as e:
            print(f"    [!] Error mengunduh {url}: {e}")

    return downloaded_paths


def generate_excel_checksheet(
    scraped_data: Dict[str, Any],
    target_part: str,
    output_excel_path: str
) -> str:
    """
    Buat file Excel checksheet (.xlsx) rapi dari data yang di-scrape dari server FactoryHub.
    """
    wb = openpyxl.Workbook()
    # Sheet 1: Cover / Info
    ws_cover = wb.active
    ws_cover.title = "Cover"

    # Sheet 2: Inspection Table
    ws_items = wb.create_sheet(title="Inspection Table")

    # Styling helpers
    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Arial", size=14, bold=True, color="1F2937")
    bold_font = Font(name="Arial", size=10, bold=True)
    regular_font = Font(name="Arial", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB")
    )

    # Populate Cover Sheet
    ws_cover.column_dimensions["A"].width = 24
    ws_cover.column_dimensions["B"].width = 50

    ws_cover["A1"] = "SUMMIT ADYAWINSA INDONESIA"
    ws_cover["A1"].font = title_font
    ws_cover["A2"] = "QUALITY CHECKSHEET MASTER (GENERATED FROM SERVER)"
    ws_cover["A2"].font = Font(name="Arial", size=11, bold=True, color="4B5563")

    info_rows = [
        ("Part Number Target", target_part),
        ("Part Number Asal", scraped_data.get("source_part", "-")),
        ("Part Name / Description", scraped_data.get("description", "-")),
        ("Document Number", scraped_data.get("doc_number", "-")),
        ("Total Inspection Points", len(scraped_data.get("items", []))),
        ("Total Reference Images", len(scraped_data.get("image_urls", []))),
        ("Waktu Duplikasi", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("URL Template Asal", scraped_data.get("edit_url", "-")),
    ]

    for row_idx, (label, val) in enumerate(info_rows, start=4):
        cell_lbl = ws_cover.cell(row=row_idx, column=1, value=label)
        cell_lbl.font = bold_font
        cell_lbl.fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
        cell_lbl.border = thin_border

        cell_val = ws_cover.cell(row=row_idx, column=2, value=str(val))
        cell_val.font = regular_font
        cell_val.border = thin_border

    # Populate Inspection Points Sheet
    headers = ["NO", "INSPECTION ITEM", "STANDARD", "METHOD", "MASTER DATA"]
    col_widths = [10, 35, 30, 25, 20]

    for col_idx, (h, w) in enumerate(zip(headers, col_widths), start=1):
        cell = ws_items.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws_items.column_dimensions[col_letter].width = w

    items = scraped_data.get("items", [])
    for row_idx, it in enumerate(items, start=2):
        row_vals = [
            it.get("item_no", "") or str(row_idx - 1),
            it.get("inspection_item", ""),
            it.get("standard", ""),
            it.get("method", ""),
            it.get("master_data", "")
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            c = ws_items.cell(row=row_idx, column=col_idx, value=val)
            c.font = regular_font
            c.border = thin_border
            if col_idx == 1:
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(vertical="center", wrap_text=True)

    os.makedirs(os.path.dirname(output_excel_path), exist_ok=True)
    wb.save(output_excel_path)
    print(f"[+] File Excel checksheet berhasil dibuat: {output_excel_path}")
    return output_excel_path


async def duplicate_server_to_server(
    source_part: str,
    target_part: str,
    submit: bool = False,
    headless: bool = False,
    browser_channel: str = "chromium",
    target_status: str = "belum"
) -> Dict[str, Any]:
    """
    Eksekusi lengkap:
    1. Scrape data part asal dari FactoryHub
    2. Unduh gambar sketsa & simpan Excel ke documents/<target_status>/<target_part>/
    3. Input template checksheet baru di FactoryHub untuk target_part
    """
    source_clean = source_part.strip().upper()
    target_clean = target_part.strip().upper()

    if not source_clean:
        raise ValueError("Nomor part sumber tidak boleh kosong.")
    if not target_clean:
        raise ValueError("Nomor part target tidak boleh kosong.")

    # Siapkan folder lokal tujuan
    safe_folder = re.sub(r'[\\/:*?"<>|]', '_', target_clean)
    target_dir = os.path.abspath(os.path.join("documents", target_status, safe_folder))
    images_dir = os.path.join(target_dir, "images")
    os.makedirs(images_dir, exist_ok=True)

    excel_filename = f"{safe_folder}_Checksheet.xlsx"
    excel_path = os.path.join(target_dir, excel_filename)

    async with async_playwright() as p:
        channel = (browser_channel or os.getenv("BROWSER_CHANNEL", "chromium")).strip().lower()
        if channel in ["edge", "ms-edge"]:
            channel = "msedge"

        browser_obj = None
        context = None
        page = None

        if channel == "chrome":
            profile_dir = os.path.expanduser("~/.factoryhub_chrome_profile")
            os.makedirs(profile_dir, exist_ok=True)
            context = await p.chromium.launch_persistent_context(
                user_data_dir=profile_dir,
                channel="chrome",
                headless=headless,
                ignore_https_errors=True,
                args=["--start-maximized", "--no-first-run", "--no-default-browser-check"],
                viewport=None if not headless else {"width": 1440, "height": 900}
            )
            page = context.pages[0] if context.pages else await context.new_page()
        else:
            browser_obj = await p.chromium.launch(
                headless=headless,
                args=["--start-maximized"]
            )
            context = await browser_obj.new_context(
                ignore_https_errors=True,
                viewport=None if not headless else {"width": 1440, "height": 900}
            )
            page = await context.new_page()

        try:
            # Step 1: Login
            await login_factoryhub(page)

            # Step 2: Scrape source template from FactoryHub
            scraped = await scrape_template_from_server(page, source_clean)

            # Step 3: Download images locally
            downloaded_imgs = []
            if scraped.get("image_urls"):
                downloaded_imgs = await download_server_images(context, scraped["image_urls"], images_dir)

            # Step 4: Generate Excel checksheet locally
            generate_excel_checksheet(scraped, target_clean, excel_path)

            # Step 5: Generate info.json
            info_path = os.path.join(target_dir, "info.json")
            info_data = {
                "part_number": target_clean,
                "part_name": scraped.get("description", ""),
                "category": "REGULAR",
                "duplicated_from_server": {
                    "source_part": source_clean,
                    "source_template_url": scraped.get("edit_url", ""),
                    "duplicated_at": datetime.now().isoformat()
                },
                "total_points": len(scraped.get("items", [])),
                "total_images": len(downloaded_imgs),
                "created_at": datetime.now().isoformat(),
                "notes": f"Diduplikasi dari server FactoryHub part {source_clean}"
            }
            with open(info_path, "w", encoding="utf-8") as f_info:
                json.dump(info_data, f_info, indent=2, ensure_ascii=False)

            clear_document_caches()

            # Step 6: Create & populate new template on FactoryHub for target_part
            print(f"\n[*] Mengisi form checksheet di FactoryHub untuk part target: {target_clean}...")
            run_result = await fill_checksheet_form(
                page=page,
                part_or_excel=target_dir,
                submit=submit,
                custom_doc_no=scraped.get("doc_number"),
                scan_images=False,  # Gunakan gambar yang baru diunduh di folder images/
                override_items=scraped.get("items")
            )

            # Step 7: Keep browser open for user review if not submitting automatically
            if not headless and not submit:
                btn = run_result.get("button_text", "Save Template")
                mode = run_result.get("mode", "CREATE")
                item_count = len(scraped.get("items", []))
                print("\n" + "="*60)
                print(f" [✓] MODE {mode}: FORM DAN {item_count} TITIK INSPEKSI TELAH TERISI LENGKAP!")
                print(" [✓] Jendela browser tetap TERBUKA di layar Anda.")
                print(f" [✓] Silakan review langsung di browser dan klik '{btn}'.")
                print(" [*] Sistem akan otomatis mendeteksi saat Anda mengklik simpan di browser...")
                print("="*60)

                start_url = page.url
                submitted_by_user = False

                if sys.stdin.isatty():
                    print("\n[*] Menunggu: Anda dapat klik tombol simpan di browser, ATAU tekan Enter di terminal ini untuk selesai...")

                    async def watch_stdin():
                        loop = asyncio.get_event_loop()
                        await loop.run_in_executor(None, input)

                    stdin_task = asyncio.create_task(watch_stdin())

                    while not stdin_task.done():
                        await asyncio.sleep(0.5)
                        try:
                            if page.is_closed():
                                print("[*] Browser ditutup oleh pengguna.")
                                break
                            curr_url = page.url
                            if curr_url != start_url and ("/create" not in curr_url and "/edit" not in curr_url):
                                print(f"\n[✓] Terdeteksi submit di browser! Form berhasil disimpan ke: {curr_url}")
                                submitted_by_user = True
                                break
                        except Exception:
                            break

                    if not stdin_task.done():
                        stdin_task.cancel()
                else:
                    print("\n[*] Menunggu interaksi Anda di browser (klik tombol simpan atau tutup browser)...")
                    for _ in range(1200):  # tunggu hingga 10 menit (setiap 500ms)
                        await asyncio.sleep(0.5)
                        try:
                            if page.is_closed():
                                print("[*] Browser ditutup oleh pengguna.")
                                break
                            curr_url = page.url
                            if curr_url != start_url and ("/create" not in curr_url and "/edit" not in curr_url):
                                print(f"\n[✓] Terdeteksi submit di browser! Form berhasil disimpan. URL: {curr_url}")
                                submitted_by_user = True
                                try:
                                    await page.wait_for_load_state("networkidle", timeout=4000)
                                except Exception:
                                    pass
                                break
                        except Exception:
                            break

                if submitted_by_user:
                    run_result["status"] = "submitted"
                    run_result["final_url"] = page.url if not page.is_closed() else ""

            # Log execution
            timestamp_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            log_line = (
                f"[{timestamp_str}] SERVER_DUPLICATE: From={source_clean} -> To={target_clean} "
                f"| Points={len(scraped.get('items', []))} | Images={len(downloaded_imgs)} "
                f"| Status={run_result.get('status', 'unknown')}\n"
            )
            try:
                with open("logs/activity.log", "a", encoding="utf-8") as lf:
                    lf.write(log_line)
            except Exception:
                pass

            return {
                "success": True,
                "source_part": source_clean,
                "target_part": target_clean,
                "target_folder": os.path.relpath(target_dir),
                "excel_path": os.path.relpath(excel_path),
                "points_count": len(scraped.get("items", [])),
                "images_count": len(downloaded_imgs),
                "form_result": run_result,
                "message": (
                    f"Berhasil menduplikasi part {source_clean} ke {target_clean} dari server FactoryHub! "
                    f"{len(scraped.get('items', []))} baris titik dan {len(downloaded_imgs)} gambar disalin."
                )
            }

        finally:
            try:
                if browser_obj:
                    await browser_obj.close()
                elif context:
                    await context.close()
            except Exception:
                pass


def main():
    import argparse
    import asyncio

    parser = argparse.ArgumentParser(
        description="Duplikasi checksheet master langsung dari server FactoryHub ke part baru."
    )
    parser.add_argument("--from", "-f", dest="src", required=True, help="Part sumber di FactoryHub (misal: 76750B040P)")
    parser.add_argument("--to", "-t", dest="target", required=True, help="Part tujuan baru (misal: 76750B040PQ)")
    parser.add_argument("--submit", action="store_true", help="Otomatis klik submit/save template di FactoryHub")
    parser.add_argument("--headless", action="store_true", help="Jalankan di background tanpa membuka jendela browser")
    parser.add_argument("--status", default="belum", choices=["belum", "done", "tidak_ada_part"], help="Kategori folder lokal")

    args = parser.parse_args()

    res = asyncio.run(duplicate_server_to_server(
        source_part=args.src,
        target_part=args.target,
        submit=args.submit,
        headless=args.headless,
        target_status=args.status
    ))

    print("\n[✓] SELESAI:")
    print(json.dumps(res, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
