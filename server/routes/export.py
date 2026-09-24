"""
Export & Google Sheets Sync API Router.
Supports 3 Worksheets: Overview, Data Master, and Log.
"""
import io
import os
from datetime import datetime
from collections import Counter
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from fastapi import APIRouter, Depends, Query, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select

from database.connection import get_db
from database.crud import list_checksheets
from database.models import Checksheet, SubmissionQueue
from services.google_sheets_service import sync_all_checksheets_to_sheet, GOOGLE_SHEET_URL

router = APIRouter(prefix="/api/export", tags=["Export"])

SPREADSHEET_URL = GOOGLE_SHEET_URL


@router.get("/excel")
async def export_excel(
    assigned_to: str = Query("ALL"),
    db: AsyncSession = Depends(get_db)
):
    """
    Generate and download styled Excel report with 3 worksheets:
    1. Overview (Ringkasan statistik & KPI)
    2. Data Master (Database semua part checksheet)
    3. Log (Histori proses, update, dan otomasi)
    """
    checksheets = await list_checksheets(session=db, assigned_to=assigned_to, limit=1000)

    # Load submission logs for the Log sheet
    res = await db.execute(
        select(SubmissionQueue).order_by(SubmissionQueue.started_at.desc()).limit(100)
    )
    submission_logs = res.scalars().all()

    wb = openpyxl.Workbook()

    # Common Styles
    indigo_fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
    indigo_light_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    section_font = Font(name="Arial", size=11, bold=True, color="1E1B4B")
    data_font = Font(name="Arial", size=9)
    bold_font = Font(name="Arial", size=9, bold=True)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # -------------------------------------------------------------
    # 1. SHEET: OVERVIEW
    # -------------------------------------------------------------
    ws_overview = wb.active
    ws_overview.title = "Overview"
    ws_overview.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_overview.merge_cells("A1:E1")
    title_cell = ws_overview.cell(1, 1, "REKAPITULASI CHECKSHEET - PT. SUMMIT ADYAWINSA INDONESIA")
    title_cell.font = Font(name="Arial", size=13, bold=True, color="FFFFFF")
    title_cell.fill = indigo_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_overview.row_dimensions[1].height = 28

    ws_overview.cell(2, 1, f"Waktu Export: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} | Filter: {assigned_to.upper()}").font = Font(name="Arial", size=9, italic=True, color="64748B")

    # Metrics Summary
    total = len(checksheets)
    done_count = sum(1 for c in checksheets if c.status == "Checksheet Done")
    ready_count = sum(1 for c in checksheets if c.status == "Belum Di Input")
    rev_count = sum(1 for c in checksheets if c.status == "Butuh Revisi")
    no_part_count = sum(1 for c in checksheets if c.status == "Tidak Ada Part")
    total_points = sum(len(c.inspection_points) if c.inspection_points else 0 for c in checksheets)

    ws_overview.cell(4, 1, "METRIK UTAMA").font = section_font
    metric_headers = ["Parameter", "Jumlah Part", "Persentase", "Keterangan"]
    for col_idx, h in enumerate(metric_headers, 1):
        c = ws_overview.cell(5, col_idx, h)
        c.font = header_font
        c.fill = indigo_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws_overview.row_dimensions[5].height = 20

    metric_rows = [
        ("Total Part Terdaftar", total, "100.0%", "Semua part dalam katalog aktif"),
        ("Checksheet Selesai (Done)", done_count, f"{(done_count / total * 100):.1f}%" if total else "0%", "Data inspeksi siap pakai"),
        ("Belum Di Input (Ready)", ready_count, f"{(ready_count / total * 100):.1f}%" if total else "0%", "Menunggu input poin/balloon"),
        ("Butuh Revisi", rev_count, f"{(rev_count / total * 100):.1f}%" if total else "0%", "Memerlukan perbaikan gambar/poin"),
        ("Tidak Ada Part", no_part_count, f"{(no_part_count / total * 100):.1f}%" if total else "0%", "Part drawing belum ditemukan"),
        ("Total Poin Inspeksi", total_points, "-", "Akumulasi seluruh item pengukuran"),
    ]
    for idx, (param, val, pct, desc) in enumerate(metric_rows, 6):
        r = [param, val, pct, desc]
        for col_idx, text in enumerate(r, 1):
            cell = ws_overview.cell(idx, col_idx, text)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in [2, 3]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # PIC Breakdown
    start_pic_row = 13
    ws_overview.cell(start_pic_row, 1, "STATISTIK PENANGGUNG JAWAB (PIC)").font = section_font
    pic_headers = ["Penanggung Jawab", "Total Part", "Done", "Belum Input", "Revisi", "Progress (%)"]
    for col_idx, h in enumerate(pic_headers, 1):
        c = ws_overview.cell(start_pic_row + 1, col_idx, h)
        c.font = header_font
        c.fill = indigo_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws_overview.row_dimensions[start_pic_row + 1].height = 20

    pic_groups = {}
    for cs in checksheets:
        pic = cs.assigned_to if (cs.assigned_to and cs.assigned_to not in ("Unassigned", "Belum Ditugaskan")) else "Belum Ditugaskan"
        pic_groups.setdefault(pic, []).append(cs)

    curr_row = start_pic_row + 2
    for pic, items in sorted(pic_groups.items(), key=lambda x: len(x[1]), reverse=True):
        p_total = len(items)
        p_done = sum(1 for c in items if c.status == "Checksheet Done")
        p_ready = sum(1 for c in items if c.status == "Belum Di Input")
        p_rev = sum(1 for c in items if c.status == "Butuh Revisi")
        p_pct = f"{(p_done / p_total * 100):.1f}%" if p_total else "0%"
        row_vals = [pic, p_total, p_done, p_ready, p_rev, p_pct]
        for col_idx, text in enumerate(row_vals, 1):
            cell = ws_overview.cell(curr_row, col_idx, text)
            cell.font = data_font
            cell.border = thin_border
            if col_idx >= 2:
                cell.alignment = Alignment(horizontal="center", vertical="center")
        curr_row += 1

    overview_widths = [28, 14, 14, 14, 14, 16]
    for i, w in enumerate(overview_widths, 1):
        col_letter = openpyxl.utils.get_column_letter(i)
        ws_overview.column_dimensions[col_letter].width = w

    # -------------------------------------------------------------
    # 2. SHEET: DATA MASTER
    # -------------------------------------------------------------
    ws_master = wb.create_sheet("Data Master")
    ws_master.views.sheetView[0].showGridLines = True

    master_headers = [
        "No",
        "Penanggung Jawab (PIC)",
        "Part Number",
        "Part Name",
        "Model",
        "Customer",
        "Doc Number",
        "Total Poin Inspeksi",
        "Status Checksheet",
        "Keterangan / Status FactoryHub",
        "Link FactoryHub",
        "Terakhir Diperbarui"
    ]
    ws_master.append(master_headers)
    ws_master.row_dimensions[1].height = 22

    for col_idx in range(1, len(master_headers) + 1):
        c = ws_master.cell(1, col_idx)
        c.fill = indigo_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")

    for idx, cs in enumerate(checksheets, 1):
        updated = cs.updated_at.strftime("%d/%m/%Y %H:%M") if cs.updated_at else "-"
        row = [
            idx,
            (cs.assigned_to if (cs.assigned_to and cs.assigned_to not in ("Unassigned", "Belum Ditugaskan")) else "Belum Ditugaskan"),
            cs.part_number or "-",
            cs.part_name or "-",
            cs.model or "-",
            cs.customer or "-",
            cs.doc_number or "-",
            len(cs.inspection_points) if cs.inspection_points else 0,
            cs.status or "-",
            cs.keterangan or "-",
            cs.factoryhub_url or "-",
            updated
        ]
        ws_master.append(row)
        for col_idx in range(1, len(row) + 1):
            cell = ws_master.cell(idx + 1, col_idx)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in [1, 2, 5, 6, 7, 8, 9, 12]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    master_widths = [6, 18, 25, 32, 12, 14, 14, 16, 18, 35, 25, 18]
    for i, w in enumerate(master_widths, 1):
        col_letter = openpyxl.utils.get_column_letter(i)
        ws_master.column_dimensions[col_letter].width = w

    # -------------------------------------------------------------
    # 3. SHEET: LOG
    # -------------------------------------------------------------
    ws_log = wb.create_sheet("Log")
    ws_log.views.sheetView[0].showGridLines = True

    log_headers = [
        "No",
        "Waktu (Timestamp)",
        "Part Number / Target",
        "Penanggung Jawab / Operator",
        "Tipe Proses / Aksi",
        "Status",
        "Keterangan / Catatan Detail"
    ]
    ws_log.append(log_headers)
    ws_log.row_dimensions[1].height = 22

    for col_idx in range(1, len(log_headers) + 1):
        c = ws_log.cell(1, col_idx)
        c.fill = indigo_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")

    log_rows = []
    # Entry 1: Export event
    log_rows.append([
        1,
        datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
        "SISTEM CHECKSHEET",
        "SYSTEM",
        "EXPORT EXCEL 3 SHEETS",
        "SUCCESS",
        f"Export data checksheet dengan total {len(checksheets)} part master"
    ])

    # Submission logs
    for sub in submission_logs:
        part_num = sub.checksheet.part_number if sub.checksheet else f"ID #{sub.checksheet_id}"
        t_time = sub.started_at.strftime("%d/%m/%Y %H:%M:%S") if sub.started_at else "-"
        detail = sub.error_message if sub.status == "FAILED" else (sub.log_output[:120] if sub.log_output else "Proses otomatisasi FactoryHub selesai")
        log_rows.append([
            len(log_rows) + 1,
            t_time,
            part_num,
            sub.operator_name or "Operator",
            "OTOMASI FACTORYHUB",
            sub.status or "-",
            detail
        ])

    # Active updates from checksheets
    active_updates = [
        cs for cs in checksheets
        if (cs.status in ("Checksheet Done", "Butuh Revisi", "Tidak Ada Part") or (cs.keterangan and cs.keterangan != "-"))
    ]
    active_updates.sort(key=lambda x: x.updated_at or datetime.min, reverse=True)

    for cs in active_updates:
        t_time = cs.updated_at.strftime("%d/%m/%Y %H:%M:%S") if cs.updated_at else "-"
        action_type = "UPDATE STATUS"
        if cs.status == "Checksheet Done":
            action_type = "INPUT SELESAI"
        elif cs.status == "Butuh Revisi":
            action_type = "PERMINTAAN REVISI"
        elif cs.status == "Tidak Ada Part":
            action_type = "PART TIDAK ADA"

        detail = cs.keterangan if (cs.keterangan and cs.keterangan != "-") else f"Pembaruan status {cs.status} dengan {len(cs.inspection_points) if cs.inspection_points else 0} poin inspeksi"
        log_rows.append([
            len(log_rows) + 1,
            t_time,
            cs.part_number,
            cs.assigned_to or "Unassigned",
            action_type,
            cs.status,
            detail
        ])

    for row_idx, r in enumerate(log_rows, 2):
        ws_log.append(r)
        for col_idx in range(1, len(r) + 1):
            cell = ws_log.cell(row_idx, col_idx)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in [1, 2, 4, 5, 6]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    log_widths = [6, 20, 25, 22, 24, 14, 50]
    for i, w in enumerate(log_widths, 1):
        col_letter = openpyxl.utils.get_column_letter(i)
        ws_log.column_dimensions[col_letter].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"REKAP_CHECKSHEET_{assigned_to.upper()}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.post("/sync-google-sheet")
@router.post("/sheet/sync")
async def sync_to_google_sheet(
    db: AsyncSession = Depends(get_db)
):
    """
    Sync all checksheets with formatted 3-worksheet table directly to the Google Spreadsheet:
    Overview, Data Master, and Log.
    """
    try:
        res = await sync_all_checksheets_to_sheet()
        if isinstance(res, dict):
            res["rows_synced"] = res.get("synced_count", 0)
        return res
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Gagal sync ke Google Spreadsheet: {str(e)}")
