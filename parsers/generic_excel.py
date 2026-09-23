"""
Generic fallback parser for Excel tabular checksheets.
Reads standard tables with Item No, Item Name, Standard, Method columns.
"""
import os
import re
from typing import Dict, List, Optional, Any
import openpyxl

from parsers.base import BaseParser
from parsers.image_extractor import extract_excel_images


class GenericExcelParser(BaseParser):
    """Generic fallback parser for any Excel-based checksheet."""

    def can_parse(self, file_path: str, wb: Optional[Any] = None) -> bool:
        return file_path.lower().endswith((".xlsx", ".xls"))

    def extract_metadata(self, file_path: str, wb: Optional[Any] = None) -> Dict[str, str]:
        doc_number = "Form 1"
        part_number = ""
        part_name = ""
        model = "-"
        customer = "-"

        should_close = False
        if wb is None:
            try:
                wb = openpyxl.load_workbook(file_path, data_only=True)
                should_close = True
            except Exception:
                wb = None

        if wb is not None:
            try:
                ws = wb.active
                for r in range(1, min(ws.max_row + 1, 15)):
                    for c in range(1, min(ws.max_column + 1, 20)):
                        val = str(ws.cell(r, c).value or "").strip()
                        val_lower = val.lower()

                        if "part no" in val_lower and not part_number:
                            if ":" in val:
                                part_number = val.split(":", 1)[1].strip()
                            else:
                                for dc in range(1, 4):
                                    cval = str(ws.cell(r, c + dc).value or "").strip()
                                    if len(cval) > 4:
                                        part_number = cval
                                        break

                        if "part name" in val_lower and not part_name:
                            if ":" in val:
                                part_name = val.split(":", 1)[1].strip()
                            else:
                                for dc in range(1, 4):
                                    cval = str(ws.cell(r, c + dc).value or "").strip()
                                    if len(cval) > 2:
                                        part_name = cval
                                        break
            finally:
                if should_close:
                    try:
                        wb.close()
                    except Exception:
                        pass

        fname = os.path.basename(file_path)
        if not part_number:
            match = re.search(r"([0-9]{4,5}[A-Z0-9_-]{3,})", fname)
            if match:
                part_number = match.group(1)

        return {
            "part_number": part_number,
            "part_name": part_name,
            "model": model,
            "customer": customer,
            "doc_number": doc_number,
            "filename": fname
        }

    def extract_inspection_points(self, file_path: str, wb: Optional[Any] = None) -> List[Dict[str, str]]:
        should_close = False
        if wb is None:
            try:
                wb = openpyxl.load_workbook(file_path, data_only=True)
                should_close = True
            except Exception:
                return []

        points = []
        try:
            ws = wb.active
            hdr_row = None
            c_no, c_item, c_std, c_tool = 1, 2, 3, 4

            for r in range(1, min(ws.max_row + 1, 25)):
                row_vals = [str(ws.cell(r, c).value or "").strip().upper() for c in range(1, 15)]
                for c, v in enumerate(row_vals, 1):
                    if v in ["NO", "NO.", "ITEM NO", "ITEM NO."]:
                        c_no = c
                        hdr_row = r
                    elif "ITEM" in v or "INSPECTION" in v:
                        c_item = c
                    elif "STANDAR" in v or "SPEC" in v or "LIMIT" in v:
                        c_std = c
                    elif "ALAT" in v or "TOOL" in v or "METHOD" in v:
                        c_tool = c
                if hdr_row:
                    break

            if not hdr_row:
                hdr_row = 1

            balloon_counter = 1
            for r in range(hdr_row + 1, min(ws.max_row + 1, hdr_row + 100)):
                v_no = str(ws.cell(r, c_no).value or "").strip()
                v_item = str(ws.cell(r, c_item).value or "").strip()
                v_std = str(ws.cell(r, c_std).value or "").strip()
                v_tool = str(ws.cell(r, c_tool).value or "").strip()

                if not v_item and not v_std:
                    continue

                if any(k in v_item.upper() for k in ["STANDAR", "INSPECTION", "POINT"]):
                    continue

                item_no = v_no if v_no and v_no.isdigit() else str(balloon_counter)
                balloon_counter += 1

                points.append({
                    "item_no": item_no,
                    "inspection_item": " ".join(v_item.split()) if v_item else f"Point {balloon_counter}",
                    "standard": " ".join(v_std.split()) if v_std else "-",
                    "method": " ".join(v_tool.split()) if v_tool else "Visual",
                    "master_data": ""
                })
        finally:
            if should_close:
                try:
                    wb.close()
                except Exception:
                    pass

        return points

    def extract_images(self, file_path: str, output_dir: Optional[str] = None, part_number: str = "") -> List[str]:
        return extract_excel_images(file_path, output_dir=output_dir, part_number=part_number)
