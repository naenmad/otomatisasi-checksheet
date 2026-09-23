"""
Parser for In-Process Quality Control (IPQC) checksheets.
Handles Appearance and Dimension multi-page inspection tables, datum markers,
and merges two-row upper/lower tolerances into canonical format (e.g. '0 + 0.3 / - 0').
"""
import os
import re
from typing import Dict, List, Optional, Any
import openpyxl

from parsers.base import BaseParser
from parsers.image_extractor import extract_excel_images


class MMKIIPQCParser(BaseParser):
    """Parser dedicated to In-Process Quality Control (IPQC) Checksheets."""

    def can_parse(self, file_path: str, wb: Optional[Any] = None) -> bool:
        if not file_path.lower().endswith((".xlsx", ".xls")):
            return False

        fp_lower = file_path.lower()
        if "ipqc" in fp_lower or "ipq" in fp_lower:
            return True

        should_close = False
        if wb is None and file_path.lower().endswith(".xlsx"):
            try:
                wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
                should_close = True
            except Exception:
                return False

        if wb is not None:
            try:
                for sname in wb.sheetnames[:5]:
                    ws = wb[sname]
                    for r in range(1, min(ws.max_row or 30, 40)):
                        c_val = str(ws.cell(r, 3).value or "").strip().upper()
                        h_val = str(ws.cell(r, 8).value or "").strip().upper()
                        if ("APPEARANCE" in c_val or "DIMENSION" in c_val) and ("STANDARD" in h_val or "METHOD" in h_val):
                            if should_close:
                                wb.close()
                            return True
            except Exception:
                pass
            finally:
                if should_close:
                    try:
                        wb.close()
                    except Exception:
                        pass

        return False

    def extract_metadata(self, file_path: str, wb: Optional[Any] = None) -> Dict[str, str]:
        doc_number = "Form 1"
        part_number = ""
        part_name = ""
        model = "-"
        customer = "PT. MMKI"

        should_close = False
        if wb is None:
            try:
                wb = openpyxl.load_workbook(file_path, data_only=True)
                should_close = True
            except Exception:
                wb = None

        if wb is not None:
            try:
                ws = wb.active
                for r in range(1, min(ws.max_row + 1, 15)):
                    for c in range(1, min(ws.max_column + 1, 20)):
                        v = str(ws.cell(r, c).value or "").strip()
                        v_u = v.upper()
                        if "PART NO" in v_u and not part_number:
                            if ":" in v:
                                part_number = v.split(":", 1)[1].strip()
                            else:
                                for dc in range(1, 4):
                                    cval = str(ws.cell(r, c + dc).value or "").strip()
                                    if len(cval) > 4:
                                        part_number = cval
                                        break
                        if "PART NAME" in v_u and not part_name:
                            if ":" in v:
                                part_name = v.split(":", 1)[1].strip()
                            else:
                                for dc in range(1, 4):
                                    cval = str(ws.cell(r, c + dc).value or "").strip()
                                    if len(cval) > 2:
                                        part_name = cval
                                        break
            finally:
                if should_close:
                    try:
                        wb.close()
                    except Exception:
                        pass

        fname = os.path.basename(file_path)
        if not part_number:
            match = re.search(r"([0-9]{4,5}[A-Z0-9_-]{3,})", fname)
            if match:
                part_number = match.group(1)

        return {
            "part_number": part_number,
            "part_name": part_name,
            "model": model,
            "customer": customer,
            "doc_number": doc_number,
            "filename": fname
        }

    def extract_inspection_points(self, file_path: str, wb: Optional[Any] = None) -> List[Dict[str, str]]:
        should_close = False
        if wb is None:
            try:
                wb = openpyxl.load_workbook(file_path, data_only=True)
                should_close = True
            except Exception:
                return []

        all_points = []
        try:
            for sname in wb.sheetnames:
                ws = wb[sname]
                is_ipqc = False
                for r in range(1, min(ws.max_row + 1, 50)):
                    c_val = str(ws.cell(r, 3).value or "").strip().upper()
                    h_val = str(ws.cell(r, 8).value or "").strip().upper()
                    l_val = str(ws.cell(r, 12).value or "").strip().upper()
                    if ("APPEARANCE" in c_val or "DIMENSION" in c_val) and ("STANDARD" in h_val or "METHOD" in l_val):
                        is_ipqc = True
                        break

                if not is_ipqc:
                    continue

                def _format_tolerance_str(s: str) -> str:
                    s = s.strip()
                    m = re.match(r"^([^\+\-\/±]+)?\s*[\+]+\s*([0-9\.]+)\s*\/\s*([\-\+])\s*([0-9\.]+)", s)
                    if m:
                        nom = (m.group(1) or "").strip()
                        u = m.group(2).strip()
                        sign2 = m.group(3).strip()
                        l = m.group(4).strip()
                        return f"{nom} + {u} / {sign2} {l}" if nom else f"+ {u} / {sign2} {l}"
                    return s

                current_section = ""
                current_item_no = ""
                current_item_name = ""
                current_method = "Visual"
                r = 1

                while r <= ws.max_row:
                    c_val = str(ws.cell(r, 3).value or "").strip()
                    h_val = str(ws.cell(r, 8).value or "").strip()
                    j_val = str(ws.cell(r, 10).value or "").strip()
                    l_val = str(ws.cell(r, 12).value or "").strip()
                    b_val = str(ws.cell(r, 2).value or "").strip()
                    a_val = str(ws.cell(r, 1).value or "").strip()

                    c_upper = c_val.upper()
                    h_upper = h_val.upper()
                    l_upper = l_val.upper()
                    a_upper = a_val.upper()

                    if "APPEARANCE" in c_upper and ("STANDARD" in h_upper or "METHOD" in l_upper):
                        current_section = "Appearance"
                        current_method = "Visual"
                        r += 1
                        continue
                    elif "DIMENSION" in c_upper and ("STANDARD" in h_upper or "METHOD" in l_upper):
                        current_section = "Dimension"
                        current_method = "Feeler Gauge"
                        r += 1
                        continue
                    elif "JUDGEMENT" in c_upper or "JML. PROD." in a_upper:
                        break

                    if current_section:
                        if any(k in c_upper for k in ["SPESIFIKASI", "WAKTU PENGECEKAN", "APPEARANCE", "DIMENSION"]) or any(k in h_upper for k in ["STANDARD", "AWAL PROSES", "AKHIR PROSES"]):
                            r += 1
                            continue

                        if b_val and (b_val.isdigit() or not any(k in b_val.upper() for k in ["NO", "ITEM"])):
                            current_item_no = b_val
                        if c_val:
                            current_item_name = c_val
                        if l_val:
                            current_method = l_val

                        datum = ""
                        if a_val and a_upper not in ["INSPECTOR", "SHIFT", "TGL", "PROSES", "NO.", "NO"]:
                            datum = a_val.strip()

                        skip_next = False
                        next_j = str(ws.cell(r + 1, 10).value or "").strip() if r + 1 <= ws.max_row else ""

                        if j_val:
                            if next_j and (next_j.startswith("-") or next_j.startswith("+")):
                                u_clean = j_val.lstrip("+").strip()
                                l_sign = "-" if next_j.startswith("-") else "+"
                                l_clean = next_j.lstrip("-+").strip()
                                std_str = f"{h_val} + {u_clean} / {l_sign} {l_clean}".strip()
                                skip_next = True
                            else:
                                std_str = f"{h_val} {j_val}".strip()
                        else:
                            std_str = h_val

                        std_formatted = _format_tolerance_str(std_str)
                        if not h_val and (std_formatted.startswith("-") or std_formatted.startswith("+")):
                            r += 1
                            continue

                        if std_formatted and std_formatted.upper() not in ["STANDARD", "AWAL PROSES", "AKHIR PROSES"] and current_item_name:
                            item_display = f"{current_item_name} [{datum}]" if datum else current_item_name
                            all_points.append({
                                "item_no": current_item_no or "1",
                                "inspection_item": " ".join(item_display.split()),
                                "standard": " ".join(std_formatted.split()),
                                "method": " ".join(current_method.split()),
                                "master_data": ""
                            })

                        if skip_next:
                            r += 1
                    r += 1
        finally:
            if should_close:
                try:
                    wb.close()
                except Exception:
                    pass

        return all_points

    def extract_images(self, file_path: str, output_dir: Optional[str] = None, part_number: str = "") -> List[str]:
        return extract_excel_images(file_path, output_dir=output_dir, part_number=part_number)
