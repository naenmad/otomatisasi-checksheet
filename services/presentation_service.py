"""
Presentation Service for Daily Report PPTX Generation.
Generates a 16:9 widescreen PowerPoint presentation for daily checksheet submissions.
- Slide 1: Executive Overview with KPI summary and operator breakdown.
- Slide 2+: Individual operator detail slides (excluding admins, operators only).
"""
import io
import re
from datetime import datetime, date, time, timedelta
from typing import List, Dict, Any, Optional

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func

from database.models import User, ActivityLog, Checksheet, InspectionPoint

# Indonesian Date Names
DAYS_ID = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"]
MONTHS_ID = [
    "", "Januari", "Februari", "Maret", "April", "Mei", "Juni",
    "Juli", "Agustus", "September", "Oktober", "November", "Desember"
]


def format_date_id(d: date) -> str:
    """Format date to Indonesian readable string: e.g. 'Senin, 28 September 2026'."""
    day_name = DAYS_ID[d.weekday()]
    month_name = MONTHS_ID[d.month]
    return f"{day_name}, {d.day} {month_name} {d.year}"


def is_operator_match(log_op: str, user: User) -> bool:
    """Match log operator string with user entity."""
    lo = (log_op or "").strip().lower()
    un = (user.username or "").strip().lower()
    nm = (user.name or "").strip().lower()
    return lo == un or lo == nm or (lo and lo in nm) or (un and un in lo)


def extract_points_from_log(
    details: str,
    cs: Optional[Checksheet],
    part_number: Optional[str],
    points_by_cs_id: Dict[int, int],
    points_by_pn: Dict[str, int]
) -> int:
    """Extract inspection points count from log details or checksheet point counts without lazy ORM IO."""
    if details:
        m = re.search(r'(\d+)\s*poin', details, re.IGNORECASE)
        if m:
            return int(m.group(1))
    if cs and cs.id in points_by_cs_id:
        return points_by_cs_id[cs.id]
    if part_number and part_number in points_by_pn:
        return points_by_pn[part_number]
    return 0


# Professional Corporate Color Palette
COLOR_NAVY = RGBColor(15, 23, 42)        # Slate 900
COLOR_HEADER_BG = RGBColor(30, 41, 59)   # Slate 800
COLOR_INDIGO = RGBColor(79, 70, 229)     # Indigo 600
COLOR_INDIGO_LIGHT = RGBColor(238, 242, 255) # Indigo 50
COLOR_EMERALD = RGBColor(16, 185, 129)   # Emerald 500
COLOR_EMERALD_LIGHT = RGBColor(236, 253, 245)
COLOR_SKY = RGBColor(14, 165, 233)       # Sky 500
COLOR_SLATE_DARK = RGBColor(30, 41, 59)  # Slate 800
COLOR_SLATE_MUTED = RGBColor(100, 116, 139) # Slate 500
COLOR_SLATE_LIGHT = RGBColor(241, 245, 249) # Slate 100
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_BORDER = RGBColor(203, 213, 225)   # Slate 300
COLOR_CARD_BG = RGBColor(255, 255, 255)


def add_slide_header(
    slide,
    badge_text: str,
    title_text: str,
    right_pill_text: str,
    right_pill_bg: RGBColor = COLOR_INDIGO
):
    """Draw a corporate header banner across top of slide."""
    # Top banner background
    banner = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0), Inches(13.333), Inches(1.15)
    )
    banner.fill.solid()
    banner.fill.fore_color.rgb = COLOR_NAVY
    banner.line.color.rgb = COLOR_NAVY

    # Left text box for badge + title
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(8.5), Inches(0.9))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p_badge = tf.paragraphs[0]
    p_badge.text = badge_text.upper()
    p_badge.font.name = "Arial"
    p_badge.font.size = Pt(8.5)
    p_badge.font.bold = True
    p_badge.font.color.rgb = RGBColor(165, 180, 252) # Indigo 300

    p_title = tf.add_paragraph()
    p_title.text = title_text
    p_title.font.name = "Arial"
    p_title.font.size = Pt(17)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_WHITE
    p_title.space_before = Pt(2)

    # Right pill badge
    if right_pill_text:
        pill = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(9.5), Inches(0.35), Inches(3.0), Inches(0.48)
        )
        pill.fill.solid()
        pill.fill.fore_color.rgb = right_pill_bg
        pill.line.color.rgb = right_pill_bg
        ptf = pill.text_frame
        ptf.vertical_anchor = MSO_ANCHOR.MIDDLE
        pp = ptf.paragraphs[0]
        pp.text = right_pill_text
        pp.alignment = PP_ALIGN.CENTER
        pp.font.name = "Arial"
        pp.font.size = Pt(10.5)
        pp.font.bold = True
        pp.font.color.rgb = COLOR_WHITE


def add_slide_footer(slide, current_slide: int, total_slides: int):
    """Draw a clean footer at bottom of slide."""
    now_str = (datetime.utcnow() + timedelta(hours=7)).strftime("%d/%m/%Y %H:%M WIB")
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.7), Inches(0.35))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"PT. Summit Adyawinsa Indonesia • Sistem Otomatisasi Checksheet QC • Digenerate: {now_str}"
    p.font.name = "Arial"
    p.font.size = Pt(8)
    p.font.color.rgb = COLOR_SLATE_MUTED

    # Page number on right
    tb_page = slide.shapes.add_textbox(Inches(11.0), Inches(7.05), Inches(1.5), Inches(0.35))
    ptf = tb_page.text_frame
    ptf.margin_left = ptf.margin_top = ptf.margin_right = ptf.margin_bottom = 0
    pp = ptf.paragraphs[0]
    pp.text = f"Slide {current_slide} / {total_slides}"
    pp.alignment = PP_ALIGN.RIGHT
    pp.font.name = "Arial"
    pp.font.size = Pt(8)
    pp.font.color.rgb = COLOR_SLATE_MUTED


async def generate_daily_report_pptx(target_date: date, session: AsyncSession) -> io.BytesIO:
    """
    Generate styled 16:9 PowerPoint report for the specified date.
    - Slide 1: Executive Overview
    - Slide 2+: Per-operator details (operators only)
    """
    # 1. Fetch only users who are operators (EXCLUDE admins)
    stmt_ops = select(User).filter(User.role == "operator").order_by(User.name.asc())
    operators: List[User] = (await session.execute(stmt_ops)).scalars().all()

    # 2. Fetch point counts pre-aggregated by checksheet and part_number to avoid lazy-loading
    stmt_pts_cs = (
        select(InspectionPoint.checksheet_id, func.count(InspectionPoint.id))
        .group_by(InspectionPoint.checksheet_id)
    )
    pts_cs_res = (await session.execute(stmt_pts_cs)).all()
    points_by_cs_id: Dict[int, int] = {row[0]: row[1] for row in pts_cs_res}

    stmt_pts_pn = (
        select(Checksheet.part_number, func.count(InspectionPoint.id))
        .join(InspectionPoint, InspectionPoint.checksheet_id == Checksheet.id)
        .group_by(Checksheet.part_number)
    )
    pts_pn_res = (await session.execute(stmt_pts_pn)).all()
    points_by_pn: Dict[str, int] = {row[0]: row[1] for row in pts_pn_res}

    # 3. Fetch all SUBMIT FACTORYHUB activity logs for this day in WIB (UTC+7)
    start_utc = datetime.combine(target_date, time.min) - timedelta(hours=7)
    end_utc = datetime.combine(target_date, time.max) - timedelta(hours=7)

    stmt_logs = (
        select(ActivityLog, Checksheet)
        .outerjoin(Checksheet, ActivityLog.part_number == Checksheet.part_number)
        .filter(ActivityLog.created_at >= start_utc, ActivityLog.created_at <= end_utc)
        .filter(ActivityLog.action == "SUBMIT FACTORYHUB")
        .order_by(ActivityLog.created_at.asc())
    )
    raw_logs = (await session.execute(stmt_logs)).all()

    # Group submissions per operator
    operator_submissions: Dict[int, List[Dict[str, Any]]] = {u.id: [] for u in operators}
    unassigned_submissions: List[Dict[str, Any]] = []

    total_submitted = len(raw_logs)
    total_points = 0
    success_count = 0

    for log, cs in raw_logs:
        pts = extract_points_from_log(log.details, cs, log.part_number, points_by_cs_id, points_by_pn)
        total_points += pts
        if log.status == "SUCCESS":
            success_count += 1

        wib_time = (log.created_at + timedelta(hours=7)).strftime("%H:%M:%S")

        sub_item = {
            "part_number": log.part_number or "-",
            "part_name": cs.part_name if cs and cs.part_name else "-",
            "model": cs.model if cs and cs.model else "-",
            "customer": cs.customer if cs and cs.customer else "-",
            "time": wib_time,
            "status": log.status or "SUCCESS",
            "points": pts,
            "details": log.details or ""
        }

        matched_user = None
        for u in operators:
            if is_operator_match(log.operator, u):
                matched_user = u
                break

        if matched_user:
            operator_submissions[matched_user.id].append(sub_item)
        else:
            unassigned_submissions.append(sub_item)

    success_rate = round((success_count / total_submitted * 100), 1) if total_submitted > 0 else 100.0
    active_operators = sum(1 for u in operators if len(operator_submissions[u.id]) > 0)
    formatted_date = format_date_id(target_date)

    # Calculate total slides count beforehand
    # Slide 1: Overview
    # Each operator: at least 1 slide. If > 10 items, ceil(items / 10) slides.
    ROWS_PER_PAGE = 10
    total_slides = 1
    operator_slide_plan = []
    for u in operators:
        subs = operator_submissions[u.id]
        if len(subs) <= ROWS_PER_PAGE:
            pages = 1
        else:
            pages = (len(subs) + ROWS_PER_PAGE - 1) // ROWS_PER_PAGE
        operator_slide_plan.append((u, subs, pages))
        total_slides += pages

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    current_slide_no = 1

    # =========================================================================
    # SLIDE 1: EXECUTIVE OVERVIEW
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    add_slide_header(
        slide1,
        badge_text="PT. SUMMIT ADYAWINSA INDONESIA • QUALITY CONTROL",
        title_text="DAILY REPORT SUBMISSION CHECKSHEET FACTORYHUB",
        right_pill_text=formatted_date,
        right_pill_bg=COLOR_INDIGO
    )

    # 4 Metric Cards
    kpis = [
        ("TOTAL CHECKSHEET SUBMIT", f"{total_submitted} Part", "Tercatat di FactoryHub hari ini", COLOR_EMERALD),
        ("TOTAL POIN INSPEKSI", f"{total_points} Poin", "Akumulasi seluruh poin ukur", COLOR_INDIGO),
        ("OPERATOR AKTIF", f"{active_operators} / {len(operators)} Orang", "Operator bertugas submit", COLOR_SKY),
        ("SUCCESS RATE", f"{success_rate}%", f"{success_count} dari {total_submitted} berhasil", COLOR_EMERALD)
    ]

    card_y = Inches(1.35)
    card_w = Inches(2.7)
    card_h = Inches(1.15)
    gap = Inches(0.3)
    start_x = Inches(0.8)

    for i, (title, val, subtext, color) in enumerate(kpis):
        cx = start_x + i * (card_w + gap)
        # Card background shape
        c_shape = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, card_w, card_h)
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = COLOR_WHITE
        c_shape.line.color.rgb = COLOR_BORDER
        c_shape.line.width = Pt(1)

        # Card text
        ctf = c_shape.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.18)
        ctf.margin_top = Inches(0.12)
        ctf.margin_right = Inches(0.18)
        ctf.margin_bottom = Inches(0.08)

        p_t = ctf.paragraphs[0]
        p_t.text = title
        p_t.font.name = "Arial"
        p_t.font.size = Pt(8.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_SLATE_MUTED

        p_v = ctf.add_paragraph()
        p_v.text = val
        p_v.font.name = "Arial"
        p_v.font.size = Pt(19)
        p_v.font.bold = True
        p_v.font.color.rgb = color
        p_v.space_before = Pt(1)

        p_s = ctf.add_paragraph()
        p_s.text = subtext
        p_s.font.name = "Arial"
        p_s.font.size = Pt(7.5)
        p_s.font.color.rgb = COLOR_SLATE_MUTED
        p_s.space_before = Pt(1)

    # Section Title: Ringkasan Kinerja Operator
    sec_tb = slide1.shapes.add_textbox(Inches(0.8), Inches(2.65), Inches(11.7), Inches(0.35))
    stf = sec_tb.text_frame
    stf.margin_left = stf.margin_top = stf.margin_right = stf.margin_bottom = 0
    sp = stf.paragraphs[0]
    sp.text = "RINGKASAN KINERJA SUBMISSION PER OPERATOR"
    sp.font.name = "Arial"
    sp.font.size = Pt(12)
    sp.font.bold = True
    sp.font.color.rgb = COLOR_SLATE_DARK

    # Summary Table per Operator
    table_rows = len(operators) + 2  # header + operators + total
    table_cols = 6
    tbl_shape = slide1.shapes.add_table(
        table_rows, table_cols,
        Inches(0.8), Inches(3.05), Inches(11.7), Inches(0.4 * table_rows)
    )
    tbl = tbl_shape.table
    col_widths = [Inches(0.6), Inches(3.2), Inches(2.0), Inches(2.0), Inches(1.9), Inches(2.0)]
    for ci, w in enumerate(col_widths):
        tbl.columns[ci].width = w

    headers = ["No", "Nama Operator", "Username / NIK", "Part Disubmit", "Estimasi Poin", "Kontribusi (%)"]
    for ci, h in enumerate(headers):
        cell = tbl.cell(0, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_HEADER_BG
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.alignment = PP_ALIGN.CENTER if ci in [0, 3, 4, 5] else PP_ALIGN.LEFT
        p.font.name = "Arial"
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE

    row_idx = 1
    total_op_parts = 0
    total_op_points = 0

    for i, u in enumerate(operators, 1):
        subs = operator_submissions[u.id]
        count_sub = len(subs)
        pts_sub = sum(s["points"] for s in subs)
        pct = (count_sub / total_submitted * 100) if total_submitted > 0 else 0.0

        total_op_parts += count_sub
        total_op_points += pts_sub

        row_data = [
            str(i),
            u.name,
            f"{u.username} ({u.nik})" if u.nik else u.username,
            f"{count_sub} Part",
            f"{pts_sub} Poin",
            f"{pct:.1f}%"
        ]

        row_bg = COLOR_WHITE if i % 2 == 1 else COLOR_SLATE_LIGHT
        if count_sub > 0:
            row_bg = COLOR_EMERALD_LIGHT if count_sub > 10 else RGBColor(240, 249, 255)

        for ci, val in enumerate(row_data):
            cell = tbl.cell(row_idx, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = row_bg
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.alignment = PP_ALIGN.CENTER if ci in [0, 3, 4, 5] else PP_ALIGN.LEFT
            p.font.name = "Arial"
            p.font.size = Pt(9.5)
            p.font.bold = (ci in [1, 3] and count_sub > 0)
            p.font.color.rgb = COLOR_SLATE_DARK

        row_idx += 1

    # Total Row
    tot_row_data = [
        "",
        "TOTAL KESELURUHAN",
        f"{len(operators)} Operator Terdaftar",
        f"{total_op_parts} Part",
        f"{total_op_points} Poin",
        "100.0%" if total_op_parts > 0 else "0.0%"
    ]
    for ci, val in enumerate(tot_row_data):
        cell = tbl.cell(row_idx, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_INDIGO_LIGHT
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.alignment = PP_ALIGN.CENTER if ci in [0, 3, 4, 5] else PP_ALIGN.LEFT
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_NAVY

    add_slide_footer(slide1, current_slide_no, total_slides)
    current_slide_no += 1

    # =========================================================================
    # SLIDE 2+: PER-OPERATOR SLIDES (EXCLUDE ADMIN, OPERATORS ONLY)
    # =========================================================================
    for u, subs, total_pages in operator_slide_plan:
        operator_part_count = len(subs)
        operator_points_count = sum(s["points"] for s in subs)

        if operator_part_count == 0:
            # 0 Submissions: Clean, modern empty-state slide
            slide = prs.slides.add_slide(blank_layout)
            add_slide_header(
                slide,
                badge_text=f"LAPORAN OPERATOR QC • PT. SUMMIT ADYAWINSA INDONESIA",
                title_text=f"DETAIL SUBMISSION: {u.name.upper()}",
                right_pill_text="0 Part Disubmit",
                right_pill_bg=COLOR_SLATE_MUTED
            )

            # Empty state container
            empty_shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(2.5), Inches(2.2), Inches(8.333), Inches(3.4)
            )
            empty_shape.fill.solid()
            empty_shape.fill.fore_color.rgb = COLOR_WHITE
            empty_shape.line.color.rgb = COLOR_BORDER
            empty_shape.line.width = Pt(1.5)

            etf = empty_shape.text_frame
            etf.word_wrap = True
            etf.vertical_anchor = MSO_ANCHOR.MIDDLE

            ep1 = etf.paragraphs[0]
            ep1.text = "BELUM ADA SUBMISSION"
            ep1.alignment = PP_ALIGN.CENTER
            ep1.font.name = "Arial"
            ep1.font.size = Pt(16)
            ep1.font.bold = True
            ep1.font.color.rgb = COLOR_SLATE_MUTED

            ep2 = etf.add_paragraph()
            ep2.text = f"Operator {u.name} belum melakukan submit checksheet ke FactoryHub pada tanggal {formatted_date}."
            ep2.alignment = PP_ALIGN.CENTER
            ep2.font.name = "Arial"
            ep2.font.size = Pt(11)
            ep2.font.color.rgb = COLOR_SLATE_MUTED
            ep2.space_before = Pt(8)

            ep3 = etf.add_paragraph()
            ep3.text = f"Username: {u.username} • NIK: {u.nik or '-'} • Status: Aktif"
            ep3.alignment = PP_ALIGN.CENTER
            ep3.font.name = "Arial"
            ep3.font.size = Pt(9.5)
            ep3.font.color.rgb = COLOR_SLATE_MUTED
            ep3.space_before = Pt(12)

            add_slide_footer(slide, current_slide_no, total_slides)
            current_slide_no += 1

        else:
            # Operator has 1 or more submissions, paginated cleanly
            for page in range(total_pages):
                slide = prs.slides.add_slide(blank_layout)
                page_info = f" (Hal {page+1}/{total_pages})" if total_pages > 1 else ""
                badge_title = f"{operator_part_count} Part Disubmit"

                add_slide_header(
                    slide,
                    badge_text=f"LAPORAN OPERATOR QC • PT. SUMMIT ADYAWINSA INDONESIA",
                    title_text=f"DETAIL SUBMISSION: {u.name.upper()}{page_info}",
                    right_pill_text=badge_title,
                    right_pill_bg=COLOR_EMERALD
                )

                # Sub-header bar
                sub_tb = slide.shapes.add_textbox(Inches(0.8), Inches(1.25), Inches(11.7), Inches(0.4))
                sutf = sub_tb.text_frame
                sutf.margin_left = sutf.margin_top = sutf.margin_right = sutf.margin_bottom = 0
                sup = sutf.paragraphs[0]
                sup.text = f"Operator: {u.name} (NIK: {u.nik or u.username})  |  Tanggal: {formatted_date}  |  Total Poin Ukur: {operator_points_count} Poin"
                sup.font.name = "Arial"
                sup.font.size = Pt(10)
                sup.font.color.rgb = COLOR_SLATE_DARK

                # Table of submissions for this page
                page_start = page * ROWS_PER_PAGE
                page_end = min(page_start + ROWS_PER_PAGE, operator_part_count)
                page_items = subs[page_start:page_end]

                table_rows = len(page_items) + 1
                table_cols = 8
                tbl_shape = slide.shapes.add_table(
                    table_rows, table_cols,
                    Inches(0.6), Inches(1.75), Inches(12.133), Inches(0.45 * table_rows)
                )
                tbl = tbl_shape.table

                # Column widths (Total = 12.133 in)
                widths = [
                    Inches(0.6),   # No
                    Inches(2.2),   # Part Number
                    Inches(3.733), # Part Name
                    Inches(1.2),   # Model
                    Inches(1.3),   # Customer
                    Inches(1.2),   # Jam Submit
                    Inches(0.9),   # Poin
                    Inches(1.0)    # Status
                ]
                for ci, w in enumerate(widths):
                    tbl.columns[ci].width = w

                col_names = ["No", "Part Number", "Nama Part", "Model", "Customer", "Jam Submit", "Poin", "Status"]
                for ci, cn in enumerate(col_names):
                    cell = tbl.cell(0, ci)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = COLOR_HEADER_BG
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    p = cell.text_frame.paragraphs[0]
                    p.text = cn
                    p.alignment = PP_ALIGN.CENTER if ci in [0, 3, 4, 5, 6, 7] else PP_ALIGN.LEFT
                    p.font.name = "Arial"
                    p.font.size = Pt(9.5)
                    p.font.bold = True
                    p.font.color.rgb = COLOR_WHITE

                for item_idx, itm in enumerate(page_items, start=page_start + 1):
                    row_i = item_idx - page_start
                    row_bg = COLOR_WHITE if row_i % 2 == 1 else COLOR_SLATE_LIGHT

                    vals = [
                        str(item_idx),
                        itm["part_number"],
                        itm["part_name"],
                        itm["model"],
                        itm["customer"],
                        itm["time"],
                        f"{itm['points']}" if itm['points'] > 0 else "-",
                        itm["status"]
                    ]

                    for ci, v in enumerate(vals):
                        cell = tbl.cell(row_i, ci)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = row_bg
                        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                        p = cell.text_frame.paragraphs[0]
                        p.text = v
                        p.alignment = PP_ALIGN.CENTER if ci in [0, 3, 4, 5, 6, 7] else PP_ALIGN.LEFT
                        p.font.name = "Arial"
                        p.font.size = Pt(9)
                        if ci == 1:
                            p.font.name = "Courier New"
                            p.font.bold = True
                        if ci == 7:
                            p.font.bold = True
                            p.font.color.rgb = COLOR_EMERALD if itm["status"] == "SUCCESS" else RGBColor(225, 29, 72)
                        else:
                            p.font.color.rgb = COLOR_SLATE_DARK

                add_slide_footer(slide, current_slide_no, total_slides)
                current_slide_no += 1

    out = io.BytesIO()
    prs.save(out)
    out.seek(0)
    return out
