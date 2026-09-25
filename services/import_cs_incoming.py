"""
Multi-Sheet Batch Importer for CS INCOMING checksheets into Supabase Database.
Extracts all parts across all sheets (including multi-sheet workbooks) for complete coverage (~360 parts).
"""
import os
import sys
import glob
import re
import asyncio
from typing import List, Dict, Any, Optional

# Ensure project root is in sys.path
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

import openpyxl
import xlrd
from sqlalchemy import delete

from database.connection import AsyncSessionLocal
from database.models import Checksheet, InspectionPoint, PartImage, SubmissionQueue
from database.crud import create_checksheet
from services.parser_service import match_catalog_status
from services.google_sheets_service import sync_all_checksheets_to_sheet
from parsers.image_extractor import extract_excel_images


def get_cell_xlsx(ws, r: int, c: int) -> str:
    v = ws.cell(r, c).value
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def get_cell_xls(sheet, r: int, c: int) -> str:
    if r - 1 < sheet.nrows and c - 1 < sheet.ncols:
        v = sheet.cell_value(r - 1, c - 1)
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v).strip()
    return ""


def extract_sheet_metadata(cell_getter, max_r: int, max_c: int, fname: str, sname: str, fp: str) -> Dict[str, str]:
    part_no = ""
    part_name = ""
    model = ""
    doc_no = "FO-45-01"

    customer = "PT. HPM"
    if "MMKI" in fp:
        customer = "PT. MMKI"
    elif "SIM" in fp:
        customer = "PT. SIM"
        if "SIM YHA" in fp:
            model = "YHA"

    for r in range(1, min(max_r + 1, 16)):
        for c in range(1, min(max_c + 1, 35)):
            v = cell_getter(r, c)
            vl = v.lower()

            if "part no" in vl and not part_no:
                if ":" in v:
                    p = v.split(":", 1)[1].strip()
                    if len(p) > 4:
                        part_no = p
                if not part_no:
                    for dc in range(1, 4):
                        cv = cell_getter(r, c + dc)
                        if len(cv) > 4 and "part no" not in cv.lower():
                            part_no = cv
                            break

            if "part name" in vl and not part_name:
                if ":" in v:
                    p = v.split(":", 1)[1].strip()
                    if len(p) > 2:
                        part_name = p
                if not part_name:
                    for dc in range(1, 4):
                        cv = cell_getter(r, c + dc)
                        if len(cv) > 2 and "part name" not in cv.lower():
                            part_name = cv
                            break

            if "model" in vl and not model:
                if ":" in v:
                    m = v.split(":", 1)[1].strip()
                    if len(m) >= 2:
                        model = m
                if not model:
                    for dc in range(1, 3):
                        cv = cell_getter(r, c + dc)
                        if len(cv) >= 2 and "model" not in cv.lower():
                            model = cv
                            break

    # Fallbacks
    if not part_no:
        clean_s = re.sub(r"\(rev\)", "", sname, flags=re.I).strip()
        if len(clean_s) >= 5 and any(ch.isdigit() for ch in clean_s):
            part_no = clean_s
        else:
            clean_fn = re.sub(r"^(CS\s*IQC\s*)", "", fname, flags=re.I)
            clean_fn = os.path.splitext(clean_fn)[0].strip()
            part_no = clean_fn

    if not model:
        if "SIM YHA" in fp:
            model = "YHA"
        else:
            parent = os.path.basename(os.path.dirname(fp))
            if parent and any(ch.isalnum() for ch in parent):
                model = parent.split(" ")[0].strip("()")

    return {
        "part_number": part_no,
        "part_name": part_name,
        "model": model or "-",
        "customer": customer,
        "doc_number": doc_no
    }


def extract_sheet_points(cell_getter, max_r: int, max_c: int) -> List[Dict[str, str]]:
    hdr_row = None
    c_no, c_item, c_std, c_tool = 1, 2, 3, 4

    for r in range(1, min(max_r + 1, 25)):
        row_vals = [cell_getter(r, c).upper() for c in range(1, min(max_c + 1, 15))]
        has_no = any(v in ["NO", "NO.", "NO / ITEM"] for v in row_vals)
        has_insp = any("INSPECTION" in v for v in row_vals)
        has_std = any("STANDAR" in v or "LIMIT" in v for v in row_vals)
        if (has_no or has_insp) and has_std:
            hdr_row = r
            for c in range(1, min(max_c + 1, 15)):
                val = cell_getter(r, c).upper()
                if val in ["NO", "NO."]:
                    c_no = c
                elif "INSPECTION" in val and "RESULT" not in val:
                    c_item = c
                elif "STANDAR" in val or "LIMIT" in val:
                    c_std = c
                elif "ALAT" in val or "TOOL" in val:
                    c_tool = c
            break

    if not hdr_row:
        return []

    balloon_counter = 1
    last_item = ""
    points = []

    for r in range(hdr_row + 1, min(max_r + 1, hdr_row + 35)):
        v_no = cell_getter(r, c_no)
        v_item = cell_getter(r, c_item)
        v_std = cell_getter(r, c_std)
        v_tool = cell_getter(r, c_tool)

        if not v_item and not v_std:
            continue
        if any(skip in v_item.upper() for skip in ["DATE", "QTY", "JUDG", "JUDGEMENT"]):
            continue
        if v_item.isdigit() and len(v_item) <= 2:
            continue

        if not v_item and v_no and not v_no.isdigit():
            v_item = v_no
            v_no = ""

        if not v_item and last_item:
            v_item = last_item
        elif v_item:
            last_item = v_item

        item_no_str = v_no if v_no else str(balloon_counter)
        if v_no and v_no.isdigit():
            balloon_counter = int(v_no) + 1
        else:
            balloon_counter += 1

        points.append({
            "item_no": item_no_str,
            "inspection_item": " ".join(v_item.split()) if v_item else f"Point {balloon_counter}",
            "standard": " ".join(v_std.split()) if v_std else "-",
            "method": " ".join(v_tool.split()) if v_tool else "Visual",
            "master_data": ""
        })

    return points


async def import_all_cs_incoming(reset_db: bool = True):
    search_pattern = "documents/CS INCOMING/**/*.xlsx"
    files = glob.glob(search_pattern, recursive=True)
    files.extend(glob.glob("documents/CS INCOMING/**/*.xls", recursive=True))
    files = [
        f for f in sorted(files)
        if not os.path.basename(f).startswith("~$")
        and "/bahan/" not in f and "\\bahan\\" not in f
    ]

    print(f"[*] Found {len(files)} CS INCOMING files to process.")

    async with AsyncSessionLocal() as session:
        if reset_db:
            print("[*] Clearing existing checksheet database records for clean import...")
            await session.execute(delete(InspectionPoint))
            await session.execute(delete(PartImage))
            await session.execute(delete(SubmissionQueue))
            await session.execute(delete(Checksheet))
            await session.commit()
            print("[✓] Database reset complete.")

        total_saved = 0
        error_count = 0
        errors: List[dict] = []

        for f_idx, fp in enumerate(files, 1):
            fname = os.path.basename(fp)
            is_xls = fp.lower().endswith(".xls")

            try:
                sheet_entries = []
                if is_xls:
                    bk = xlrd.open_workbook(fp)
                    for sname in bk.sheet_names():
                        ws = bk.sheet_by_name(sname)
                        meta = extract_sheet_metadata(
                            cell_getter=lambda r, c, s=ws: get_cell_xls(s, r, c),
                            max_r=ws.nrows,
                            max_c=ws.ncols,
                            fname=fname,
                            sname=sname,
                            fp=fp
                        )
                        pts = extract_sheet_points(
                            cell_getter=lambda r, c, s=ws: get_cell_xls(s, r, c),
                            max_r=ws.nrows,
                            max_c=ws.ncols
                        )
                        sheet_entries.append({
                            "sheet": sname,
                            "meta": meta,
                            "points": pts,
                            "images": []
                        })
                else:
                    wb = openpyxl.load_workbook(fp, data_only=True)
                    for sname in wb.sheetnames:
                        ws = wb[sname]
                        meta = extract_sheet_metadata(
                            cell_getter=lambda r, c, s=ws: get_cell_xlsx(s, r, c),
                            max_r=ws.max_row or 15,
                            max_c=ws.max_column or 35,
                            fname=fname,
                            sname=sname,
                            fp=fp
                        )
                        pts = extract_sheet_points(
                            cell_getter=lambda r, c, s=ws: get_cell_xlsx(s, r, c),
                            max_r=ws.max_row or 15,
                            max_c=ws.max_column or 35
                        )
                        # Extract images for part
                        imgs = []
                        try:
                            imgs = extract_excel_images(fp, part_number=meta["part_number"])
                        except Exception:
                            pass

                        sheet_entries.append({
                            "sheet": sname,
                            "meta": meta,
                            "points": pts,
                            "images": imgs
                        })
                    wb.close()

                # Deduplicate rev sheets within the same file for the same part number
                deduped = {}
                for entry in sheet_entries:
                    pno = entry["meta"]["part_number"]
                    norm_pn = re.sub(r"[^a-zA-Z0-9]", "", pno).upper()
                    key = (fname, norm_pn)
                    if key in deduped:
                        # If duplicate part number in same file, prefer (rev) sheet or sheet with more points
                        if "rev" in entry["sheet"].lower() or len(entry["points"]) > len(deduped[key]["points"]):
                            deduped[key] = entry
                    else:
                        deduped[key] = entry

                # Save each distinct sheet/part to database
                for key, item in deduped.items():
                    meta = item["meta"]
                    part_no = meta["part_number"]
                    pts = item["points"]
                    imgs = item["images"]

                    cat_match = match_catalog_status(part_no)

                    await create_checksheet(
                        session=session,
                        part_number=part_no,
                        part_name=meta["part_name"],
                        model=meta["model"],
                        customer=meta["customer"],
                        doc_number=meta["doc_number"],
                        template_type="IQCIncomingParser",
                        status=cat_match["status"],
                        assigned_to="Unassigned",
                        keterangan=cat_match["keterangan"],
                        raw_file_path=fp,
                        points=pts,
                        images=imgs
                    )
                    total_saved += 1

                if f_idx % 15 == 0 or f_idx == len(files):
                    print(f"[{f_idx}/{len(files)} files processed] -> {total_saved} checksheets created in DB")

            except Exception as e:
                error_count += 1
                errors.append({"file": fname, "error": str(e)})
                print(f"[!] Error on {fname}: {e}")

    print("\n" + "=" * 55)
    print(f"[+] Multi-Sheet Batch Import Complete!")
    print(f"    Total Excel Files : {len(files)}")
    print(f"    Total Checksheets : {total_saved}")
    print(f"    Errors Encountered: {error_count}")
    print("=" * 55)

    print("\n[*] Starting automatic Google Sheets sync for all imported checksheets...")
    try:
        await sync_all_checksheets_to_sheet()
        print("[✓] Google Sheets synchronization finished successfully!")
    except Exception as e:
        print(f"[!] Warning: Google Sheets sync failed: {e}")

    return {
        "files_processed": len(files),
        "checksheets_saved": total_saved,
        "errors": error_count,
        "error_details": errors
    }


if __name__ == "__main__":
    asyncio.run(import_all_cs_incoming(reset_db=True))
