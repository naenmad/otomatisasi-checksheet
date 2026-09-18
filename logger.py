"""
Logging Module for Checksheet Master Automation
Manages:
- Excel dual-sheet logbook: logs/history.xlsx ("Execution History" & "Search History")
- Structured JSON log: logs/execution_history.json
- Text activity log: logs/activity.log
"""

import os
import json
from datetime import datetime
from typing import Dict, List, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

LOGS_DIR = os.path.abspath("logs")
EXCEL_LOG_PATH = os.path.join(LOGS_DIR, "history.xlsx")
JSON_LOG_PATH = os.path.join(LOGS_DIR, "execution_history.json")
TEXT_LOG_PATH = os.path.join(LOGS_DIR, "activity.log")


def ensure_logs_dir():
    os.makedirs(LOGS_DIR, exist_ok=True)


def _init_excel_workbook() -> openpyxl.Workbook:
    """Initialize Excel history workbook with formatted header rows."""
    ensure_logs_dir()
    if os.path.isfile(EXCEL_LOG_PATH):
        try:
            return openpyxl.load_workbook(EXCEL_LOG_PATH)
        except Exception:
            pass

    wb = openpyxl.Workbook()
    # Sheet 1: Execution History
    ws_exec = wb.active
    ws_exec.title = "Execution History"

    exec_headers = [
        "Timestamp", "Part Number", "Document File", "Type",
        "Mode Scan", "Result Action", "Points Count", "Images Count",
        "Diff Summary", "Notes"
    ]
    ws_exec.append(exec_headers)

    # Sheet 2: Search History
    ws_search = wb.create_sheet(title="Search History")
    search_headers = [
        "Timestamp", "Search Query", "Part Number", "FH Status",
        "Category", "Template ID", "Local Document", "Notes"
    ]
    ws_search.append(search_headers)

    # Style headers
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for ws, headers in [(ws_exec, exec_headers), (ws_search, search_headers)]:
        ws.row_dimensions[1].height = 26
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(1, col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = 20

    ws_exec.column_dimensions["A"].width = 20  # Timestamp
    ws_exec.column_dimensions["B"].width = 24  # Part Number
    ws_exec.column_dimensions["C"].width = 32  # Document File
    ws_exec.column_dimensions["I"].width = 45  # Diff Summary
    ws_exec.column_dimensions["J"].width = 30  # Notes

    ws_search.column_dimensions["A"].width = 20  # Timestamp
    ws_search.column_dimensions["B"].width = 25  # Search Query
    ws_search.column_dimensions["C"].width = 25  # Part Number
    ws_search.column_dimensions["D"].width = 16  # FH Status
    ws_search.column_dimensions["E"].width = 22  # Category
    ws_search.column_dimensions["G"].width = 20  # Local Document

    wb.save(EXCEL_LOG_PATH)
    return wb


def log_execution(
    part_number: str,
    doc_file: str,
    file_type: str,
    scan_mode: str,
    result_action: str,
    points_count: int,
    images_count: int,
    diff_summary: str = "-",
    diff_details: Optional[Dict[str, Any]] = None,
    notes: str = ""
):
    """
    Log an automation run (Create / Edit / Error) to Excel, JSON, and text log.
    """
    ensure_logs_dir()
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # 1. Append to Excel
    try:
        wb = _init_excel_workbook()
        if "Execution History" in wb.sheetnames:
            ws = wb["Execution History"]
        else:
            ws = wb.create_sheet(title="Execution History")

        row = [
            now_str,
            part_number,
            os.path.basename(doc_file),
            file_type.upper(),
            scan_mode,
            result_action.upper(),
            points_count,
            images_count,
            diff_summary,
            notes
        ]
        ws.append(row)

        # Style result action cell
        last_row = ws.max_row
        res_cell = ws.cell(last_row, 6)
        if "BARU" in result_action.upper():
            res_cell.font = Font(bold=True, color="0070C0")
        elif "UPDATE" in result_action.upper():
            res_cell.font = Font(bold=True, color="ED7D31")
        elif "ERROR" in result_action.upper() or "TIDAK" in result_action.upper():
            res_cell.font = Font(bold=True, color="C00000")

        wb.save(EXCEL_LOG_PATH)
    except Exception as e:
        print(f"[Warning] Failed to write Excel execution log: {e}")

    # 2. Append to JSON
    try:
        data = []
        if os.path.isfile(JSON_LOG_PATH):
            try:
                with open(JSON_LOG_PATH, "r", encoding="utf-8") as f:
                    data = json.load(f)
            except Exception:
                data = []

        entry = {
            "timestamp": now_str,
            "part_number": part_number,
            "doc_file": os.path.basename(doc_file),
            "file_path": doc_file,
            "file_type": file_type,
            "scan_mode": scan_mode,
            "result_action": result_action.upper(),
            "points_count": points_count,
            "images_count": images_count,
            "diff_summary": diff_summary,
            "diff_details": diff_details or {},
            "notes": notes
        }
        data.insert(0, entry)  # Prepend newest

        # Keep max 500 entries
        data = data[:500]
        with open(JSON_LOG_PATH, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False)
    except Exception as e:
        print(f"[Warning] Failed to write JSON execution log: {e}")

    # 3. Append to Text log
    try:
        with open(TEXT_LOG_PATH, "a", encoding="utf-8") as f:
            f.write(f"[{now_str}] RUN: Part={part_number} | Mode={scan_mode} | Action={result_action.upper()} | Points={points_count} | Images={images_count} | Diff={diff_summary} | Notes={notes}\n")
    except Exception:
        pass


def log_search(
    search_query: str,
    part_number: str,
    fh_status: str,
    category: str = "-",
    template_id: str = "-",
    local_doc: str = "-",
    notes: str = ""
):
    """
    Log a search query result to Excel and text log.
    """
    ensure_logs_dir()
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    try:
        wb = _init_excel_workbook()
        if "Search History" in wb.sheetnames:
            ws = wb["Search History"]
        else:
            ws = wb.create_sheet(title="Search History")

        row = [
            now_str,
            search_query,
            part_number,
            fh_status.upper(),
            category,
            template_id,
            local_doc,
            notes
        ]
        ws.append(row)

        last_row = ws.max_row
        status_cell = ws.cell(last_row, 4)
        if "ADA" in fh_status.upper() and "TIDAK" not in fh_status.upper():
            status_cell.font = Font(bold=True, color="385723")
        else:
            status_cell.font = Font(bold=True, color="C00000")

        wb.save(EXCEL_LOG_PATH)
    except Exception as e:
        print(f"[Warning] Failed to write Excel search log: {e}")

    try:
        with open(TEXT_LOG_PATH, "a", encoding="utf-8") as f:
            f.write(f"[{now_str}] SEARCH: Query='{search_query}' | Part={part_number} | Status={fh_status.upper()} | Cat={category} | Local={local_doc}\n")
    except Exception:
        pass


def get_recent_execution_logs(limit: int = 50) -> List[Dict[str, Any]]:
    """Retrieve recent execution logs from JSON."""
    if not os.path.isfile(JSON_LOG_PATH):
        return []
    try:
        with open(JSON_LOG_PATH, "r", encoding="utf-8") as f:
            data = json.load(f)
            return data[:limit]
    except Exception:
        return []
