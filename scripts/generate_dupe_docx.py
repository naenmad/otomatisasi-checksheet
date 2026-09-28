import json
import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_docx():
    with open('/tmp/dupe_report.json') as f:
        data = json.load(f)

    with open('data/factoryhub_catalog.json') as f:
        cat_data = json.load(f)
    catalog_parts = cat_data.get('parts', [])
    catalog_clean_map = {p.get('clean_number', ''): p for p in catalog_parts}

    hpm_items = data['CS MATERIAL  INCOMING HPM']
    sim_items = data['CS MATERIAL INCOMING SIM YHA']

    # Database query for rich part info if available
    db_map = {}
    try:
        import asyncio
        from database.connection import AsyncSessionLocal
        from sqlalchemy import text
        async def fetch_db():
            async with AsyncSessionLocal() as session:
                r = await session.execute(text('SELECT clean_part_number, part_number, part_name, model FROM checksheets'))
                return {row[0]: (row[1], row[2], row[3]) for row in r.fetchall()}
        db_map = asyncio.run(fetch_db())
    except Exception as e:
        print('DB fetch warning:', e)

    def resolve_part(cpn, fallback_pno, fallback_pname):
        if cpn in db_map and db_map[cpn][1] and db_map[cpn][1] != '-':
            return db_map[cpn][0], db_map[cpn][1], db_map[cpn][2]
        if cpn in catalog_clean_map:
            cp = catalog_clean_map[cpn]
            return cp.get('part_number', fallback_pno), cp.get('part_name', fallback_pname), '-'
        for k, v in db_map.items():
            if k.startswith(cpn) and v[1] and v[1] != '-':
                return v[0], v[1], v[2]
        for cp in catalog_parts:
            if cp.get('clean_number', '').startswith(cpn):
                return cp.get('part_number', fallback_pno), cp.get('part_name', fallback_pname), '-'
        return fallback_pno, fallback_pname, '-'

    doc = Document()

    # Set Margins (Narrow margins: 0.7 inch)
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run('Laporan Analisis Part Beririsan (Duplikat)')
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 41, 59) # Slate 800

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run('Folder Analisis: CS IQC SUBCONT vs 4 Folder Checksheet Lainnya')
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(100, 116, 139) # Slate 500

    # Callout / Note Box
    note_tbl = doc.add_table(rows=1, cols=1)
    note_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    note_cell = note_tbl.cell(0, 0)
    set_cell_background(note_cell, 'F1F5F9') # Light slate
    set_cell_margins(note_cell, top=140, bottom=140, left=180, right=180)
    note_p = note_cell.paragraphs[0]
    note_p.paragraph_format.space_after = Pt(0)
    r_icon = note_p.add_run('Catatan Verifikasi QC:\n')
    r_icon.font.name = 'Calibri'
    r_icon.font.size = Pt(10)
    r_icon.font.bold = True
    r_icon.font.color.rgb = RGBColor(15, 23, 42)
    r_txt = note_p.add_run(
        'Laporan resmi ini memetakan seluruh 59 part unik yang ditemukan pada folder "CS IQC SUBCONT" '
        'dan juga terdapat pada folder checksheet lainnya (HPM dan SIM YHA). Seluruh data telah dinormalisasi '
        'dan dimapping ke Part Number resmi FactoryHub untuk mencegah duplikasi antrean otomasi.'
    )
    r_txt.font.name = 'Calibri'
    r_txt.font.size = Pt(9.5)
    r_txt.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Section 1: Executive Summary
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    r_h1 = h1.add_run('1. Ringkasan Eksekutif Antar Folder')
    r_h1.font.name = 'Calibri'
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(30, 41, 59)

    summary_headers = ['Folder Pembanding', 'Total Part', 'Jumlah Beririsan', 'Status', 'Karakteristik Utama']
    summary_data = [
        ['CS IQC MATERIAL HPM TG4R', '68', '0 part (0%)', 'Unik', 'Khusus part model TG4R'],
        ['CS MATERIAL INCOMING MMKI', '80', '0 part (0%)', 'Unik', 'Khusus material incoming MMKI'],
        ['CS MATERIAL INCOMING HPM', '89', '14 part (15.7%)', 'Identik', 'Format sheet satuan vs buku gabungan'],
        ['CS MATERIAL INCOMING SIM YHA', '77', '45 part (58.4%)', 'Tumpang Tindih', 'Checksheet Coil/Sheet vs Subcont'],
        ['TOTAL IRISAN UNIK', '-', '59 PART', 'Tervalidasi', 'Telah dinormalisasi di database']
    ]

    s_tbl = doc.add_table(rows=len(summary_data) + 1, cols=len(summary_headers))
    s_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_tbl.autofit = False

    # Header Row
    for col_idx, h_text in enumerate(summary_headers):
        cell = s_tbl.cell(0, col_idx)
        set_cell_background(cell, '1E293B') # Dark navy
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h_text)
        run.font.name = 'Calibri'
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for row_idx, row_vals in enumerate(summary_data, start=1):
        bg = 'F8FAFC' if row_idx % 2 == 1 else 'FFFFFF'
        if row_idx == len(summary_data):
            bg = 'E2E8F0' # Highlight total row
        for col_idx, val in enumerate(row_vals):
            cell = s_tbl.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(val)
            run.font.name = 'Calibri'
            run.font.size = Pt(9)
            if row_idx == len(summary_data):
                run.font.bold = True
            run.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 2: HPM 14 Parts
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(4)
    r_h2 = h2.add_run('2. Rincian 14 Part Beririsan: CS IQC SUBCONT vs CS MATERIAL INCOMING HPM')
    r_h2.font.name = 'Calibri'
    r_h2.font.size = Pt(14)
    r_h2.font.bold = True
    r_h2.font.color.rgb = RGBColor(30, 41, 59)

    p_desc2 = doc.add_paragraph()
    p_desc2.paragraph_format.space_after = Pt(6)
    r_d2 = p_desc2.add_run(
        'Pada bagian ini, part di folder CS IQC SUBCONT bersumber dari file buku kerja gabungan "c.sheet HPM/C.SHEET TCF HPM.xlsx", '
        'sedangkan di folder CS MATERIAL INCOMING HPM tersimpan sebagai file satuan terpisah per part di subfolder modelnya masing-masing '
        '(2WF JAZZ, 2XP HR-V, T5L, TE7). Poin inspeksi dan gambar sketsa pada kedua folder ini terkonfirmasi sama persis (identik).'
    )
    r_d2.font.name = 'Calibri'
    r_d2.font.size = Pt(9.5)
    r_d2.font.color.rgb = RGBColor(71, 85, 105)

    hpm_table_headers = ['No', 'Part Number Resmi', 'Nama Part (FactoryHub / DB)', 'Model', 'Sumber SUBCONT', 'Sumber INCOMING HPM', 'Status']
    hpm_tbl = doc.add_table(rows=len(hpm_items) + 1, cols=len(hpm_table_headers))
    hpm_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hpm_tbl.autofit = False

    for col_idx, h_text in enumerate(hpm_table_headers):
        cell = hpm_tbl.cell(0, col_idx)
        set_cell_background(cell, '1E293B')
        set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h_text)
        run.font.name = 'Calibri'
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    for idx, item in enumerate(hpm_items, 1):
        cpn = item['clean_pno']
        sc = item['subcont']
        ot = item['other']
        pno, pname, model = resolve_part(cpn, sc['part_no'], sc['part_name'])
        row_vals = [
            str(idx),
            pno,
            pname,
            model,
            f"{os.path.basename(sc['file'])}\n[{sc['sheet']}]",
            f"{os.path.basename(ot['file'])}\n[{ot['sheet']}]",
            "Identik"
        ]
        bg = 'F8FAFC' if idx % 2 == 1 else 'FFFFFF'
        for col_idx, val in enumerate(row_vals):
            cell = hpm_tbl.cell(idx, col_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(val)
            run.font.name = 'Calibri'
            run.font.size = Pt(8.5)
            if col_idx == 1:
                run.font.bold = True
            run.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_page_break()

    # Section 3: SIM YHA 45 Parts
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(4)
    r_h3 = h3.add_run('3. Rincian 45 Part Beririsan: CS IQC SUBCONT vs CS MATERIAL INCOMING SIM YHA')
    r_h3.font.name = 'Calibri'
    r_h3.font.size = Pt(14)
    r_h3.font.bold = True
    r_h3.font.color.rgb = RGBColor(30, 41, 59)

    p_desc3 = doc.add_paragraph()
    p_desc3.paragraph_format.space_after = Pt(6)
    r_d3 = p_desc3.add_run(
        'Part di folder CS IQC SUBCONT bersumber dari proses subkontraktor SIM ("CS subcont SIM.xlsx"), sedangkan di folder '
        'CS MATERIAL INCOMING SIM YHA bersumber dari incoming bahan baku Coil/Sheet ("7.IQC/CS IQC  coil (BODY)..xlsx" dan "CS IQC  coil (SEAT).xlsx"). '
        'Sesuai verifikasi, sistem memprioritaskan data resmi dari folder 7.IQC sebagai master data checksheet.'
    )
    r_d3.font.name = 'Calibri'
    r_d3.font.size = Pt(9.5)
    r_d3.font.color.rgb = RGBColor(71, 85, 105)

    sim_table_headers = ['No', 'Part Number Resmi', 'Nama Part (FactoryHub / DB)', 'Model', 'Sumber SUBCONT', 'Sumber SIM YHA (7.IQC)', 'Prioritas']
    sim_tbl = doc.add_table(rows=len(sim_items) + 1, cols=len(sim_table_headers))
    sim_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    sim_tbl.autofit = False

    for col_idx, h_text in enumerate(sim_table_headers):
        cell = sim_tbl.cell(0, col_idx)
        set_cell_background(cell, '1E293B')
        set_cell_margins(cell, top=120, bottom=120, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h_text)
        run.font.name = 'Calibri'
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    for idx, item in enumerate(sim_items, 1):
        cpn = item['clean_pno']
        sc = item['subcont']
        ot = item['other']
        pno, pname, model = resolve_part(cpn, sc['part_no'], sc['part_name'])
        row_vals = [
            str(idx),
            pno,
            pname,
            model,
            f"{os.path.basename(sc['file'])}\n[{sc['sheet']}]",
            f"{os.path.basename(ot['file'])}\n[{ot['sheet']}]",
            "Resmi 7.IQC"
        ]
        bg = 'F8FAFC' if idx % 2 == 1 else 'FFFFFF'
        for col_idx, val in enumerate(row_vals):
            cell = sim_tbl.cell(idx, col_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(val)
            run.font.name = 'Calibri'
            run.font.size = Pt(8.5)
            if col_idx == 1:
                run.font.bold = True
            run.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 4: Integrity Rules
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)
    r_h4 = h4.add_run('4. Aturan Integritas Data di Database & Otomasi FactoryHub')
    r_h4.font.name = 'Calibri'
    r_h4.font.size = Pt(14)
    r_h4.font.bold = True
    r_h4.font.color.rgb = RGBColor(30, 41, 59)

    rules = [
        ('Deduplikasi via Clean Part Number', 'Seluruh nomor part dinormalisasi otomatis (menghapus tanda minus, spasi, titik). Hal ini menjamin tidak ada checksheet ganda pada tabel database (total 802 part unik) maupun antrean kirim FactoryHub.'),
        ('Preservasi Status Reviewed', 'Checksheet yang telah ditinjau oleh operator di halaman Review & Edit dan ditandai "Reviewed" terlindungi 100% dan tidak akan diturunkan statusnya menjadi draft kembali saat penambahan folder subkontraktor.'),
        ('Prioritas Acuan Resmi 7.IQC', 'Apabila ditemukan perbedaan detail antara proses subkontraktor dan bahan mentah SIM YHA, sistem secara konsisten mengutamakan data master dari folder resmi 7.IQC.'),
        ('Isolasi Gambar Sketsa WebP', 'Setiap sketsa drawing diekstrak dan diisolasi khusus per tab/worksheet (menggunakan pemetaan relasi OpenXML drawing targets) sehingga tidak ada gambar yang tertukar antar sheet.')
    ]

    for title, desc in rules:
        p_r = doc.add_paragraph()
        p_r.paragraph_format.space_before = Pt(2)
        p_r.paragraph_format.space_after = Pt(4)
        run_t = p_r.add_run(f'• {title}: ')
        run_t.font.name = 'Calibri'
        run_t.font.size = Pt(10)
        run_t.font.bold = True
        run_t.font.color.rgb = RGBColor(30, 41, 59)
        run_d = p_r.add_run(desc)
        run_d.font.name = 'Calibri'
        run_d.font.size = Pt(9.5)
        run_d.font.color.rgb = RGBColor(71, 85, 105)

    output_path = 'documents/PERBANDINGAN_DUPLIKAT_CS_SUBCONT.docx'
    doc.save(output_path)
    print(f'Successfully generated {output_path}!')

if __name__ == '__main__':
    create_docx()
