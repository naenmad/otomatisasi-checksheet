#!/usr/bin/env python3
"""
Update subcont checksheets from their original Excel workbooks.
Extracts inspection points, standardizes via TextNormalizer, extracts drawings,
and populates the database.

Usage:
  python scripts/update_subcont_checksheets.py --file "documents/CS INCOMING/CS IQC SUBCONT/C.sheet MMKI/AAP- MMKI.xlsx"
  python scripts/update_subcont_checksheets.py --file "documents/CS INCOMING/CS IQC SUBCONT/C.sheet MMKI/AAP- MMKI.xlsx" --apply
  python scripts/update_subcont_checksheets.py --all-empty --apply
"""
import os
import sys
import glob
import zipfile
import asyncio
import argparse
import openpyxl

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from database.connection import AsyncSessionLocal
from database.crud import create_checksheet
from database.models import Checksheet, InspectionPoint, PartImage
from sqlalchemy import select, delete, text
from parsers.smart_parser import TextNormalizer, FuzzyToolNormalizer
from parsers.image_extractor import get_workbook_sheet_drawing_map, extract_sheet_images, extract_excel_images
from scripts.import_cs_subcont import (
    parse_subcont_sheet_meta,
    parse_subcont_sheet_points,
    clean_str,
    DUMMY_SHEETS
)
from services.parser_service import match_catalog_status


async def process_file(fp: str, session, cs_by_pn: dict, apply_mode: bool = False):
    fname = os.path.basename(fp)
    if not os.path.exists(fp):
        print(f"[!] File not found: {fp}")
        return 0, 0, 0

    print(f"\nProcessing workbook: {fp}")
    z_archive = zipfile.ZipFile(fp, "r") if zipfile.is_zipfile(fp) else None
    drawing_map = get_workbook_sheet_drawing_map(z_archive) if z_archive else {}

    try:
        wb = openpyxl.load_workbook(fp, data_only=True)
    except Exception as e:
        print(f"[!] Failed to load workbook {fp}: {e}")
        if z_archive:
            z_archive.close()
        return 0, 0, 0

    updated_count = 0
    created_count = 0
    total_pts_count = 0

    for sname in wb.sheetnames:
        sname_clean = sname.strip().lower()
        if sname_clean in DUMMY_SHEETS or not sname_clean:
            continue

        ws = wb[sname]
        matrix = list(ws.iter_rows(values_only=True, max_row=60, max_col=35))
        meta = parse_subcont_sheet_meta(matrix, fp, sname)
        part_no = meta["part_number"]
        cpn = clean_str(part_no)

        if not cpn or len(cpn) < 3 or cpn in DUMMY_SHEETS:
            continue

        raw_pts = parse_subcont_sheet_points(matrix)
        pts = TextNormalizer.expand_points(raw_pts)
        if not pts:
            continue

        imgs = []
        try:
            d_target = drawing_map.get(sname_clean)
            imgs = extract_sheet_images(fp, sheet_name=sname, part_number=part_no, z=z_archive, drawing_target=d_target)
            if not imgs and len(wb.sheetnames) == 1:
                imgs = extract_excel_images(fp, part_number=part_no)
        except Exception as e:
            pass

        cat_match = match_catalog_status(part_no)
        existing_cs = cs_by_pn.get(cpn)
        action = "UPDATE" if existing_cs else "CREATE"

        print(f"  [{action}] Sheet [{sname:15s}] -> PN: {part_no:20s} | Name: {meta['part_name']:25s} | Pts: {len(pts):2d} | Imgs: {len(imgs)}")

        total_pts_count += len(pts)
        if existing_cs:
            updated_count += 1
        else:
            created_count += 1

        if apply_mode:
            if existing_cs:
                existing_cs.part_name = meta["part_name"] or existing_cs.part_name
                existing_cs.raw_file_path = fp
                if meta.get("customer") and meta["customer"] != "PT. SUMMIT":
                    existing_cs.customer = meta["customer"]

                # Replace inspection points
                await session.execute(delete(InspectionPoint).where(InspectionPoint.checksheet_id == existing_cs.id))
                for idx, pt in enumerate(pts):
                    ip = InspectionPoint(
                        checksheet_id=existing_cs.id,
                        item_no=pt.get("item_no") or str(idx + 1),
                        inspection_item=pt.get("inspection_item") or f"Point {idx + 1}",
                        standard=pt.get("standard") or "-",
                        method=pt.get("method") or "Visual",
                        master_data=pt.get("master_data") or "",
                        order_index=idx
                    )
                    session.add(ip)

                # Replace images if new images found
                if imgs:
                    await session.execute(delete(PartImage).where(PartImage.checksheet_id == existing_cs.id))
                    for img_p in imgs:
                        pimg = PartImage(
                            checksheet_id=existing_cs.id,
                            image_path=img_p,
                            image_url=img_p
                        )
                        session.add(pimg)
            else:
                new_cs = await create_checksheet(
                    session=session,
                    part_number=part_no,
                    part_name=meta["part_name"],
                    model=meta["model"],
                    customer=meta["customer"],
                    doc_number=meta["doc_number"],
                    template_type="IQCSubcontParser",
                    status=cat_match["status"],
                    assigned_to="Unassigned",
                    keterangan=cat_match["keterangan"] or f"Diimpor dari {fname}",
                    raw_file_path=fp,
                    points=pts,
                    images=imgs
                )
                cs_by_pn[cpn] = new_cs

    wb.close()
    if z_archive:
        z_archive.close()

    return created_count, updated_count, total_pts_count


async def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", type=str, help="Specific Excel file to process")
    parser.add_argument("--all-empty", action="store_true", help="Process all files that have empty checksheets")
    parser.add_argument("--apply", action="store_true", help="Commit changes to database")
    args = parser.parse_args()

    apply_mode = args.apply

    async with AsyncSessionLocal() as session:
        # Load all existing checksheets map
        res = await session.execute(select(Checksheet))
        existing_all = res.scalars().all()
        cs_by_pn = {}
        for cs in existing_all:
            cpn = clean_str(cs.part_number)
            if cpn:
                cs_by_pn[cpn] = cs

        files_to_process = []
        if args.file:
            files_to_process.append(args.file)
        elif args.all_empty:
            # Find distinct raw_file_path for 0-point checksheets
            res = await session.execute(text('''
                SELECT DISTINCT c.raw_file_path
                FROM checksheets c
                WHERE (SELECT COUNT(*) FROM inspection_points ip WHERE ip.checksheet_id = c.id) = 0
                  AND c.raw_file_path IS NOT NULL AND c.raw_file_path != ''
                ORDER BY c.raw_file_path;
            '''))
            files_to_process = [r[0] for r in res.fetchall()]
        else:
            # Default to AAP- MMKI.xlsx (which contains 73234B010P)
            files_to_process.append("documents/CS INCOMING/CS IQC SUBCONT/C.sheet MMKI/AAP- MMKI.xlsx")

        print(f"Total workbooks to process: {len(files_to_process)}")
        print(f"Mode: {'APPLY' if apply_mode else 'DRY-RUN'}")
        print("=" * 80)

        tot_created = 0
        tot_updated = 0
        tot_points = 0

        for fp in files_to_process:
            c, u, p = await process_file(fp, session, cs_by_pn, apply_mode)
            tot_created += c
            tot_updated += u
            tot_points += p

        if apply_mode:
            await session.commit()
            print("\nCHANGES COMMITTED TO DATABASE!")
        else:
            print("\nDRY-RUN COMPLETE (no changes written). Run with --apply to commit.")

        print(f"\nFinal Summary:")
        print(f"  Created: {tot_created}")
        print(f"  Updated: {tot_updated}")
        print(f"  Total points: {tot_points}")


if __name__ == "__main__":
    asyncio.run(main())
