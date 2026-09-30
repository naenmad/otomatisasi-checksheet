#!/usr/bin/env python3
"""
Populate and Standardize Inspection Points for Need Review ("Butuh Revisi") Checksheets.
========================================================================================
Extracts table data directly from raw Excel workbooks across diverse layouts without image scanning.
Applies all project standardization rules:
  1. Burry: standard ≤ 0.3 mm, method Caliper
  2. Appearance & Painting defect items: No Rust, No Scratch, No Dent, No Crack, No Wave,
     No Wrinkle, No Neck, No Over Cutting, No Spatter, No Bubble, No Meler, No Kotor,
     No Bintik Putih, No Menggumpal, No Orange Peel, No Cacat, Profile OK, Packing.
     All have standard "OK / NG" (Profile OK: "Sesuai Sample"), method "Visual".
  3. One Numbering Rule: Consecutive defect rows share the exact same item_no.
  4. General tool & nominal standardizations.

Usage:
    python scripts/populate_need_review_checksheets.py            # Preview (dry-run)
    python scripts/populate_need_review_checksheets.py --apply    # Commit to DB
"""
import asyncio
import os
import sys
import re
from datetime import datetime
from typing import List, Dict, Any, Optional

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

import openpyxl
from sqlalchemy import select, delete, func
from database.connection import AsyncSessionLocal
from database.models import Checksheet, InspectionPoint, ActivityLog
from parsers.smart_parser import TextNormalizer, FuzzyToolNormalizer


def find_matching_sheet(wb, cs: Checksheet) -> Optional[str]:
    """Find the sheet corresponding to checksheet in workbook."""
    clean_pn = re.sub(r"[^A-Za-z0-9]", "", cs.part_number).upper()
    pn_lower = cs.part_number.lower()
    name_lower = (cs.part_name or "").lower()

    # 1. Exact match with sheet name
    for s in wb.sheetnames:
        clean_s = re.sub(r"[^A-Za-z0-9]", "", s).upper()
        if clean_s == clean_pn:
            return s

    # 2. Substring match
    for s in wb.sheetnames:
        clean_s = re.sub(r"[^A-Za-z0-9]", "", s).upper()
        if len(clean_s) >= 4 and clean_s in clean_pn:
            return s
        if len(clean_pn) >= 4 and clean_pn in clean_s:
            return s

    # 3. Known special mappings
    if "weld nut m5" in name_lower or "mf927220" in pn_lower:
        for candidate in ["m5", "m5 (hpm", "m5 sl", "m5 (2)"]:
            if candidate in wb.sheetnames:
                return candidate
    elif "weld nut m8" in name_lower or "mf927420" in pn_lower:
        if "m8" in wb.sheetnames:
            return "m8"
    elif "weld nut m6" in name_lower or "mf927320" in pn_lower:
        for candidate in ["m6", "m6 (2)", "m6 (3)", "m6 SL"]:
            if candidate in wb.sheetnames:
                return candidate
    elif "mu000366" in pn_lower and "0336" in wb.sheetnames:
        return "0336"
    elif "mf198034" in pn_lower and "8034" in wb.sheetnames:
        return "8034"
    elif "mf198052" in pn_lower:
        for candidate in ["1420", "198052", "1420 (2)"]:
            if candidate in wb.sheetnames:
                return candidate
    elif "mf927522" in pn_lower:
        for candidate in ["hex M10", "1420 (2)", "90306"]:
            if candidate in wb.sheetnames:
                return candidate
    elif "6722a771" in pn_lower and "62133" in wb.sheetnames:
        return "62133"
    elif "61771-bz160" in pn_lower or "61771-bz150" in pn_lower:
        for candidate in ["61771-BZ150-60", "61771-BZ150", "61771-BZ160"]:
            if candidate in wb.sheetnames:
                return candidate
    elif "61772-bz140" in pn_lower or "61772-bz150" in pn_lower:
        for candidate in ["61772-BZ140-50", "61772-BZ140", "61772-BZ150"]:
            if candidate in wb.sheetnames:
                return candidate
    elif "mb137096" in pn_lower and "RAIL" in wb.sheetnames:
        return "RAIL"
    elif "mb137171x" in pn_lower and "RAIL 1" in wb.sheetnames:
        return "RAIL 1"
    elif "71250-3m0" in pn_lower:
        for candidate in ["t86", "Sheet2", "t86 Ex"]:
            if candidate in wb.sheetnames:
                return candidate
    elif "mb032883" in pn_lower:
        for candidate in ["17528", "7529", "rail"]:
            if candidate in wb.sheetnames:
                return candidate

    # 4. Search part number inside cell contents of sheets (scans up to column 30)
    for s in wb.sheetnames:
        ws = wb[s]
        for r in range(1, 16):
            for c in range(1, 31):
                val = str(ws.cell(r, c).value or "").upper()
                if cs.clean_part_number in val or (len(clean_pn) >= 5 and clean_pn in re.sub(r"[^A-Za-z0-9]", "", val)):
                    return s

    if len(wb.sheetnames) == 1:
        return wb.sheetnames[0]

    return None


def extract_table_points_from_sheet(ws) -> List[Dict[str, str]]:
    """Extract inspection points table from worksheet handling diverse column layouts."""
    # 1. Locate header row
    hdr_row = None
    c_no, c_item, c_std, c_tool = 1, 2, 3, 4

    for r in range(1, min(ws.max_row + 1, 25)):
        row_vals = [str(ws.cell(r, c).value or "").strip().upper() for c in range(1, 12)]
        has_no = any(v in ["NO", "NO.", "NO / ITEM", "ITEM"] for v in row_vals)
        has_insp = any("INSPECT" in v or "POINT" in v for v in row_vals)
        has_std = any("STANDAR" in v or "LIMIT" in v or ("SPEC" in v and "INSPEC" not in v) for v in row_vals)

        if (has_no or has_insp) and has_std:
            hdr_row = r
            for c in range(1, 12):
                val = str(ws.cell(r, c).value or "").strip().upper()
                if val in ["NO", "NO.", "NO / ITEM"]:
                    c_no = c
                elif "INSPECT" in val and "RESULT" not in val:
                    c_item = c
                elif ("STANDAR" in val or "LIMIT" in val or ("SPEC" in val and "INSPEC" not in val)) and "RESULT" not in val:
                    c_std = c
                elif ("ALAT" in val or "TOOL" in val or "METODE" in val) and "RESULT" not in val:
                    c_tool = c
            break

    if not hdr_row:
        # Fallback search for default inspection table header
        for r in range(10, 20):
            val1 = str(ws.cell(r, 1).value or "").strip().upper()
            val2 = str(ws.cell(r, 2).value or "").strip().upper()
            if "NO" in val1 or "INSPECTION" in val2:
                hdr_row = r
                break

    if not hdr_row:
        return []

    # Read data rows
    raw_points = []
    balloon_counter = 1
    last_item = ""

    for r in range(hdr_row + 1, min(ws.max_row + 1, hdr_row + 35)):
        row_str_all = " ".join(str(ws.cell(r, c).value or "") for c in range(1, 15)).upper()

        # Stop if footer or conformity section reached
        stop_keywords = ["CONFORMITY", "GOOD FOR USE", "REJECT", "DIBUAT", "DIPERIKSA", "NO.DOC", "REASON", "CATATAN"]
        if any(kw in row_str_all for kw in stop_keywords):
            break

        # Filter out sample/header repetitions (e.g. 1 2 3 4 5, Date :, Qty :)
        if any(skip in row_str_all for skip in ["DATE :", "QTY :", "ARRIVAL DATE", "CHECK DATE", "JUDGEMENT"]):
            continue

        v_no = str(ws.cell(r, c_no).value or "").strip()
        v_item = str(ws.cell(r, c_item).value or "").strip()
        v_std = str(ws.cell(r, c_std).value or "").strip()
        v_tool = str(ws.cell(r, c_tool).value or "").strip()

        # Skip noise row if col 1..3 are sample numbers
        if v_no in ["1", "2", "3", "4", "5"] and v_item in ["2", "3", "4", "5"] and v_std in ["3", "4", "5"]:
            continue
        if v_no in ["1", "2", "3", "4", "5"] and not v_item and not v_std:
            continue

        if not v_item and not v_std:
            continue

        # If v_no is actually item name
        if not v_item and v_no and not v_no.isdigit():
            v_item = v_no
            v_no = ""

        # Forward-fill grouped item
        if not v_item and last_item:
            v_item = last_item
        elif v_item:
            last_item = v_item

        item_no_str = v_no if v_no else str(balloon_counter)
        if v_no and v_no.isdigit():
            balloon_counter = int(v_no) + 1
        else:
            balloon_counter += 1

        raw_points.append({
            "item_no": item_no_str,
            "inspection_item": " ".join(v_item.split()) if v_item else f"Point {balloon_counter}",
            "standard": " ".join(v_std.split()) if v_std else "-",
            "method": " ".join(v_tool.split()) if v_tool else "Visual",
            "master_data": ""
        })

    return raw_points


async def run_population(apply_mode: bool = False):
    print("=" * 75)
    print("[*] POPULATE & STANDARDIZE NEED REVIEW ('BUTUH REVISI') CHECKSHEETS")
    print(f"[*] Mode: {'>>> APPLY (COMMITTING TO DATABASE) <<<' if apply_mode else 'DRY-RUN (PREVIEW ONLY)'}")
    print("=" * 75)

    async with AsyncSessionLocal() as session:
        # Fetch checksheets with status 'Butuh Revisi'
        res = await session.execute(
            select(Checksheet)
            .filter(Checksheet.status == "Butuh Revisi")
            .order_by(Checksheet.id.asc())
        )
        checksheets = res.scalars().all()
        print(f"[*] Menemukan {len(checksheets)} checksheet berstatus 'Butuh Revisi'.")

        wb_cache = {}
        success_count = 0
        skipped_count = 0
        total_points_added = 0

        for cs in checksheets:
            fp = cs.raw_file_path
            if not fp or not os.path.exists(fp):
                print(f"[!] CS #{cs.id} ({cs.part_number}): File tidak ditemukan ({fp})")
                skipped_count += 1
                continue

            if fp not in wb_cache:
                try:
                    wb_cache[fp] = openpyxl.load_workbook(fp, data_only=True)
                except Exception as e:
                    print(f"[!] CS #{cs.id} ({cs.part_number}): Gagal membuka workbook ({e})")
                    skipped_count += 1
                    continue

            wb = wb_cache[fp]
            matched_sheet = find_matching_sheet(wb, cs)
            if not matched_sheet:
                print(f"[?] CS #{cs.id} ({cs.part_number}): Tidak ditemukan sheet yang cocok di {os.path.basename(fp)}")
                skipped_count += 1
                continue

            ws = wb[matched_sheet]
            raw_pts = extract_table_points_from_sheet(ws)
            if not raw_pts:
                print(f"[-] CS #{cs.id} ({cs.part_number}): Sheet '{matched_sheet}' tidak memiliki baris inspeksi terbaca")
                skipped_count += 1
                continue

            # Standardize & expand points
            standardized_pts = TextNormalizer.expand_points(raw_pts)
            success_count += 1
            total_points_added += len(standardized_pts)

            if success_count <= 8 or success_count % 25 == 0:
                print(f"\n[CS #{cs.id}] {cs.part_number} ({cs.part_name}) -> Sheet: '{matched_sheet}' ({len(standardized_pts)} poin)")
                for p in standardized_pts[:5]:
                    print(f"    #{p['item_no']:2s} | {p['inspection_item']:22s} | {p['standard']:18s} | {p['method']}")
                if len(standardized_pts) > 5:
                    print(f"    ... (+{len(standardized_pts) - 5} baris lainnya)")

            if apply_mode:
                # Delete old points
                await session.execute(
                    delete(InspectionPoint).where(InspectionPoint.checksheet_id == cs.id)
                )
                # Insert standardized points
                for idx, pt in enumerate(standardized_pts):
                    ip = InspectionPoint(
                        checksheet_id=cs.id,
                        item_no=pt.get("item_no") or str(idx + 1),
                        inspection_item=pt.get("inspection_item") or f"Point {idx + 1}",
                        standard=pt.get("standard") or "-",
                        method=pt.get("method") or "Visual",
                        master_data=pt.get("master_data") or "",
                        order_index=idx
                    )
                    session.add(ip)

        if apply_mode:
            log = ActivityLog(
                action="REPARSE_NEED_REVIEW",
                operator="System Population",
                status="SUCCESS",
                details=f"Reparsed and standardized {success_count} Need Review checksheets ({total_points_added} points created)",
                created_at=datetime.utcnow()
            )
            session.add(log)
            await session.commit()
            print("\n" + "=" * 75)
            print("[✓] PENAMBAHAN DATA TABEL SELESAI & TERSIMPAN DI DATABASE!")
        else:
            print("\n" + "=" * 75)
            print("[*] DRY-RUN SELESAI (Database belum diubah).")

        print(f"    - Checksheet Butuh Revisi      : {len(checksheets)}")
        print(f"    - Berhasil Diekstrak & Standar : {success_count}")
        print(f"    - Dilewati / Belum Ada Sheet   : {skipped_count}")
        print(f"    - Total Titik Inspeksi Dibuat  : {total_points_added}")
        print("=" * 75)


if __name__ == "__main__":
    apply_mode_flag = "--apply" in sys.argv
    asyncio.run(run_population(apply_mode=apply_mode_flag))
